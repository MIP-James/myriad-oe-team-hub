r"""MFi Report 엑셀 → mfi_brands.local.json (자동 수집 제외 브랜드 목록).

  py -3.13 build_mfi.py "…\MFi Report 10_19_26.xlsx"

브랜드사 제공 자료라 결과 파일은 gitignore (공개 레포 금지).
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OUT = HERE / "mfi_brands.local.json"

# 브랜드명이 이 단어들과 같으면 제목에 흔히 나오는 일반어라 오탐 → 목록에서 뺌
GENERIC = {
    "apple", "iphone", "ipad", "ipod", "mac", "macbook", "watch", "airpods", "magsafe", "lightning",
    "case", "cases", "cover", "cable", "cables", "charger", "charging", "adapter", "power", "usb", "type-c",
    "pro", "max", "mini", "plus", "ultra", "air", "one", "red", "blue", "black", "white", "pink", "gold",
    "smart", "premium", "design", "classic", "simple", "basic", "original", "oem", "new", "best", "home",
    "mobile", "phone", "audio", "sound", "music", "car", "auto", "led", "light", "go", "on", "it", "in",
    "n/a", "none", "na", "tbd", "generic", "brand", "no brand", "unbranded", "accessory", "accessories",
    "dock", "speaker", "speakers", "headphones", "earphones", "band", "strap", "glass", "film", "magnetic",
    # 제목에선 '어느 기기용'을 뜻하는 기기 제조사·모델명
    "google", "pixel", "samsung", "galaxy", "lg", "sony", "huawei", "xiaomi", "microsoft", "motorola", "nokia",
    # 케이스 제목에 흔한 일반어
    "uni", "cec", "rugged", "outdoor", "clear", "crystal", "hybrid", "armor", "armour", "slim", "tough",
    "shield", "defender", "guard", "silicone", "leather", "wallet", "stand", "grip", "ring", "fusion",
    "edge", "flip", "fold", "duo", "neo", "lite", "pure", "matte", "glossy", "frame", "bumper", "military",
}
SUFFIX = re.compile(r"[,.]?\s*(co\.?,?\s*ltd|ltd|limited|inc|llc|corp(oration)?|gmbh|s\.?a|b\.?v|plc|pty|co)\.?$", re.I)


def clean(name):
    n = re.sub(r"\s+", " ", (name or "").strip().lower())
    n = SUFFIX.sub("", n).strip(" ,.")
    return n


def main():
    src = sys.argv[1]
    ws = load_workbook(src, read_only=True).worksheets[0]
    brands = set()
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        for part in str(row[3] or "").split(","):
            b = clean(part)
            if len(b) >= 3 and b not in GENERIC and not b.isdigit():
                brands.add(b)
    OUT.write_text(json.dumps({"source": Path(src).name, "brands": sorted(brands)}, ensure_ascii=False, indent=0),
                   encoding="utf-8")
    print(f"{len(brands)} brands → {OUT}")


if __name__ == "__main__":
    main()
