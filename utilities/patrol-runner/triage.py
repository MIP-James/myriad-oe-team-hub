"""순찰 결과 후처리 — 신규 URL 을 A(바로 BPM) / B(사람 검토) / C(참고) / X(제외) 로 나눈다.

브랜드별 단어 목록·기준값은 설정 파일(keywords.local.json 의 "triage")에 둔다 (공개 레포 금지).
판정 자체는 BPM·M-Verify 가 하고, 여기서는 "무엇부터 볼지"와 "확실히 안 볼 것"만 정한다.

  py -3.13 triage.py output\\patrol_2026-10-27_0936.xlsx      # 기존 순찰 엑셀 재분류
"""
import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent

DEFAULTS = {
    "brand_terms": [],          # 제목에 있으면 관련 상품 (예: 브랜드명·제품 계열명)
    "brand_strong_terms": [],   # 브랜드를 직접 가리키는 말 — 정품 표방은 이게 함께 있을 때만 인정
    "genuine_claim_terms": [],  # '정품으로 판매' 신호 (상표 경로)
    "claim_window": 4,          # 표방어가 브랜드 앞뒤 몇 글자 안에 있어야 하는지
    "official_line_terms": [],  # 브랜드 공식 제품 라인명 — 제목 어디에 있든 정품 표방으로 봄
    "compat_terms": ["호환"],   # '호환품으로 판매' 신호 (디자인 경로)
    "noise_terms": [],          # 거의 항상 대상 아님 (액세서리 부속 등) → C
    "seller_min_listings": 0,   # 같은 판매자 리스팅이 이보다 적으면 one-off 의심 (0 = 사용 안 함)
    "low_price_ratio": 0.4,     # 가격 상한 대비 이 비율 이하면 저가 신호
    "tier_a": 70,               # 정품 표방 또는 재침해자여야 도달
    "tier_b": 35,               # 관련어 + 판매자 규모 정상이면 B (사람이 디자인 확인)
    "official_sellers": {},     # {"플랫폼명": ["스토어명", ...]} — 2단계에서 BPM client/list 로 대체
    "known_sellers": {},        # {"플랫폼명": ["스토어명", ...]} — 재침해자(타겟 보드) 2단계에서 DB 로 대체
    "bpm_registered_urls": [],  # 2단계에서 BPM sell-info 조회로 대체
}


def _has(text, terms):
    t = (text or "").lower()
    return [w for w in terms if w and w.lower() in t]


def _near(text, anchors, terms, window):
    """anchors 중 하나의 앞뒤 window 글자 안에 있는 terms 를 돌려준다."""
    t = (text or "").lower()
    found = []
    for a in anchors:
        for m in re.finditer(re.escape(a.lower()), t):
            around = t[max(0, m.start() - window): m.end() + window]
            found += [w for w in terms if w.lower() in around and w not in found]
    return found


def _norm_url(url):
    u = re.sub(r"^https?://(www\.)?", "", (url or "").strip().lower())
    return u.split("?")[0].rstrip("/")


def load_rules(cfg):
    rules = dict(DEFAULTS)
    rules.update((cfg or {}).get("triage", {}))
    return rules


def triage(rows, keywords_by_text, rules, seller_counts=None):
    """rows: 순찰 신규 행(dict). 각 행에 tier/score/reasons 를 붙여 돌려준다."""
    if seller_counts is None:
        seller_counts = Counter((r["platform"], r.get("seller") or "") for r in rows)
    official = {p: {s.lower() for s in v} for p, v in rules["official_sellers"].items()}
    known = {p: {s.lower() for s in v} for p, v in rules["known_sellers"].items()}
    registered = {_norm_url(u) for u in rules["bpm_registered_urls"]}

    for r in rows:
        title, seller, plat = r.get("title") or "", (r.get("seller") or "").strip(), r["platform"]
        kw = keywords_by_text.get(r.get("keyword"), {})
        reasons, score, tier = [], 30, None

        # --- 확실히 제외 (X)
        if seller.lower() in official.get(plat, set()):
            tier, reasons = "X", ["공식 판매처"]
        elif _norm_url(r.get("url")) in registered:
            tier, reasons = "X", ["BPM 등록됨"]

        if tier is None:
            # --- 관련도
            brand = _has(title, rules["brand_terms"])
            if brand:
                score += 15
                reasons.append(f"관련어: {brand[0]}")
            else:
                score -= 20
                reasons.append("관련어 없음")
            # --- 판매 방식: 정품 표방(상표 경로) vs 호환 표기(디자인 경로)
            # 호환품은 디자인 모방 여부로 갈리므로 제목만으론 판단 불가 → 점수 없이 근거만 남김 (B 이하)
            # '정품 스누피'·'매장정품 (타 브랜드)' 같은 오탐 방지: 브랜드 바로 옆(앞뒤 window 글자)의 표방어만 인정
            claim = (_has(title, rules["official_line_terms"])
                     or _near(title, rules["brand_strong_terms"], rules["genuine_claim_terms"], rules["claim_window"]))
            if claim:
                score += 30
                reasons.append(f"정품 표방: {claim[0]}")
                # 가격은 '정품으로 판매' 리스팅에서만 의미 있음
                cap, price = kw.get("max_price"), r.get("price")
                if cap and price and price <= cap * rules["low_price_ratio"]:
                    score += 10
                    reasons.append(f"정품 대비 저가 ({price:,}원)")
            elif _has(title, rules["compat_terms"]):
                reasons.append("호환 표기 (디자인 확인 필요)")
            # --- 판매자
            if seller.lower() in known.get(plat, set()):
                score += 30
                reasons.append("재침해자")
            n = seller_counts.get((plat, seller), 0)
            if rules["seller_min_listings"]:
                if n >= rules["seller_min_listings"]:
                    score += 10
                    reasons.append(f"판매자 리스팅 {n}건")
                else:
                    score -= 15
                    reasons.append(f"one-off 의심 (리스팅 {n}건)")
            # --- 잡음
            noise = _has(title, rules["noise_terms"])
            if noise:
                score -= 25
                reasons.append(f"부속·잡음: {noise[0]}")
            tier = "A" if score >= rules["tier_a"] else "B" if score >= rules["tier_b"] else "C"
            # 판매자 규모는 오늘 결과만 보고 센 값이라 one-off 판단이 약함 → A 까지는 못 가게만
            if tier == "A" and any(x.startswith("one-off") for x in reasons):
                tier = "B"

        r.update(tier=tier, score=score if tier != "X" else 0, reasons=" · ".join(reasons))
    rows.sort(key=lambda r: ("ABCX".index(r["tier"]), -r["score"]))
    return rows


def _read_new_rows(xlsx):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True).worksheets[0]
    keys = {"플랫폼": "platform", "키워드": "keyword", "예상 Model Type": "model_type", "상품명": "title",
            "가격(원)": "price", "판매자": "seller", "상품 URL": "url", "광고": "ad"}
    it = ws.iter_rows(values_only=True)
    head = [keys.get(h, h) for h in next(it)]
    return [dict(zip(head, row)) for row in it if any(row)]


def main():
    ap = argparse.ArgumentParser(description="순찰 결과 등급 분류")
    ap.add_argument("xlsx")
    ap.add_argument("--config", default=str(HERE / "keywords.local.json"))
    args = ap.parse_args()
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    rows = _read_new_rows(args.xlsx)
    rows = triage(rows, {k["keyword"]: k for k in cfg["keywords"]}, load_rules(cfg))
    import report
    out = Path(args.xlsx).with_name(Path(args.xlsx).stem + "_triage.xlsx")
    report.write_triage(out, rows)
    c = Counter(r["tier"] for r in rows)
    print(f"A {c['A']} / B {c['B']} / C {c['C']} / X {c['X']}  → {out}")


if __name__ == "__main__":
    sys.exit(main())
