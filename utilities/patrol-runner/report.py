"""자동 수집 결과 엑셀 (신규 / 제외 / 실행 요약)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD_FILL = PatternFill("solid", fgColor="3A3737")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="B98A00", underline="single")
LINE = Side(style="thin", color="E7E3DE")
BORDER = Border(top=LINE, bottom=LINE, left=LINE, right=LINE)

ITEM_COLS = [
    ("플랫폼", "platform", 10), ("키워드", "keyword", 22), ("예상 Model Type", "model_type", 22),
    ("상품명", "title", 60), ("가격(원)", "price", 11), ("판매자", "seller", 18),
    ("상품 URL", "url", 55), ("광고", "ad", 6), ("썸네일", "image", 40),
]
SUMMARY_COLS = [
    ("플랫폼", "platform", 10), ("키워드", "keyword", 26), ("본 상품", "found", 9),
    ("신규(포함)", "new", 10), ("메모", "note", 40),
]


def _sheet(ws, cols, rows):
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        vals = []
        for _, key, _w in cols:
            v = r.get(key, "")
            if key == "ad":
                v = "광고" if v else ""
            vals.append(v)
        ws.append(vals)
        row = ws.max_row
        for i, (_, key, _w) in enumerate(cols, start=1):
            c = ws.cell(row=row, column=i)
            c.border = BORDER
            c.alignment = Alignment(vertical="center")
            if key == "price" and isinstance(c.value, int):
                c.number_format = "#,##0"
            if key in ("url", "image") and c.value:
                c.hyperlink, c.font = c.value, LINK_FONT
    for i, (_, _k, w) in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


TRIAGE_COLS = [("등급", "tier", 6), ("점수", "score", 6), ("근거", "reasons", 40),
               ("이미지 거리", "img_dist", 9)] + ITEM_COLS
TIER_SHEETS = [("A", "대표 이미지 일치"), ("B", "검토"), ("C", "참고"), ("X", "제외")]


def _triage_sheets(wb, rows):
    for tier, title in TIER_SHEETS:
        part = [r for r in rows if r.get("tier") == tier]
        _sheet(wb.create_sheet(f"{title} {len(part)}건"), TRIAGE_COLS, part)


def write_triage(path, rows):
    wb = Workbook()
    wb.remove(wb.active)
    _triage_sheets(wb, rows)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write(path, new_rows, excluded_rows, summary, started, finished):
    if new_rows and "tier" in new_rows[0]:
        wb = Workbook()
        wb.remove(wb.active)
        _triage_sheets(wb, new_rows)
        _sheet(wb.create_sheet(f"필터 제외 {len(excluded_rows)}건"), ITEM_COLS + [("제외 사유", "reason", 30)], excluded_rows)
        _sheet(wb.create_sheet("실행 요약"), SUMMARY_COLS, summary)
        wb["실행 요약"].cell(row=len(summary) + 3, column=1, value=f"자동 수집 {started:%Y-%m-%d %H:%M} ~ {finished:%H:%M}")
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = f"신규 {len(new_rows)}건"
    _sheet(ws, ITEM_COLS, new_rows)
    _sheet(wb.create_sheet(f"제외 {len(excluded_rows)}건"), ITEM_COLS + [("제외 사유", "reason", 30)], excluded_rows)
    ws3 = wb.create_sheet("실행 요약")
    _sheet(ws3, SUMMARY_COLS, summary)
    ws3.cell(row=ws3.max_row + 2, column=1, value=f"자동 수집 {started:%Y-%m-%d %H:%M} ~ {finished:%H:%M}")
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
