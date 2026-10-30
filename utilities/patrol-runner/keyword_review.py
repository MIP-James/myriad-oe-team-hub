"""검색 키워드 점검 리포트 — 허브 판정 결과로 키워드별 적중률·제외어 후보·새 키워드 후보를 뽑는다.

허브 '자동 수집 결과'에서 실무자가 판정한 상품(patrol_items)을 받아:
  확정 = BPM 전송 완료(sent) 또는 BPM 등록됨 / 상세 확인 = review / 제외 = 사람이 '아님' 처리한 것
  (자동 제외·미판정은 적중률 계산에서 뺌, 판매자 매장 순회 결과도 키워드 검색이 아니라서 뺌)
시트:
  1. 키워드 성적표 — 키워드 × 플랫폼 적중률 + 판단 (끄기·좁히기·늘리기 후보)
  2. 제외어 후보 — 제외 상품명에만 자꾸 나오는 단어 (Model Type 별)
  3. 새 키워드 후보 — 확정·상세 확인 상품명에 자꾸 나오는데 지금 키워드에 없는 단어·두 단어 묶음
  4. 가격 상한 — 확정 상품 가격 분포 ↔ 설정 max_price
  5. 확정·상세 확인 상품명 — 원문 (사람이 직접 읽고 판단)
  6. 현재 키워드 설정

  py -3.13 keyword_review.py                    # 전체 기간
  py -3.13 keyword_review.py --since 2026-11-01 # 이 날 이후 처음 본 상품만
키 = admin-scripts/.env 의 SUPABASE_SERVICE_ROLE_KEY (읽기만 함). 결과 = output/keyword_review_날짜.xlsx
"""
import argparse
import json
import os
import re
import sys
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import median

from openpyxl import Workbook

from report import _sheet

HERE = Path(__file__).resolve().parent
ENV_PATH = HERE.parent.parent / "admin-scripts" / ".env"
SUPABASE_URL = "https://sweaecmkofxyunnwpjfh.supabase.co"
COLS = "platform,product_id,url,title,price,seller,keyword,model_type,status,excluded_by,bpm_registered,reasons,first_seen"
STORE_VIA = "판매자 매장 순회"

MIN_JUDGED = 20       # 판정 이만큼은 쌓여야 키워드 판단
TERM_MIN_NEG = 8      # 제외어 후보: 제외 상품명 최소 등장 수
TERM_MIN_POS = 3      # 새 키워드 후보: 확정·상세 확인 상품명 최소 등장 수
TERM_LIFT = 1.5       # 새 키워드 후보: 그 Model Type 평균 적중률의 이 배 이상 (원래 잘 맞는 유형에서 흔한 단어 거름)
TOP_EXCL, TOP_NEW = 20, 15  # Model Type 당 후보 수
_TOKEN = re.compile(r"[0-9a-zA-Z가-힣]+")
STOP = {"to", "for", "the", "and", "with", "of", "in", "용", "및", "외", "등"}  # 뜻 없는 연결어


def service_key():
    k = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if k:
        return k
    if ENV_PATH.exists():
        for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
            if line.strip().startswith("SUPABASE_SERVICE_ROLE_KEY="):
                return line.split("=", 1)[1].strip().strip('"').strip("'")
    sys.exit(f"Supabase 키를 찾지 못했습니다: {ENV_PATH} 의 SUPABASE_SERVICE_ROLE_KEY")


def fetch(key, since=None):
    rows, step = [], 1000
    flt = f"&first_seen=gte.{since}" if since else ""
    while True:
        req = urllib.request.Request(
            f"{SUPABASE_URL}/rest/v1/patrol_items?select={COLS}{flt}&order=id&offset={len(rows)}&limit={step}",
            headers={"apikey": key, "Authorization": f"Bearer {key}", "User-Agent": "MyriadAutoCollect/1.0"})
        with urllib.request.urlopen(req, timeout=60) as resp:
            part = json.loads(resp.read())
        rows += part
        if len(part) < step:
            return rows


def label(r):
    """확정 / 상세 확인 / 제외 / '' (미판정·자동 제외)."""
    if r["status"] == "sent" or r["bpm_registered"]:
        return "확정"
    if r["status"] == "review":
        return "상세 확인"
    if r["status"] == "excluded" and r.get("excluded_by"):
        return "제외"
    return ""


def tokens(title):
    """상품명 → 단어 집합 (소문자, 숫자만인 것 제외) + 붙어 있는 두 단어 묶음."""
    ws = [w.lower() for w in _TOKEN.findall(title or "") if len(w) >= 2 and not w.isdigit() and w.lower() not in STOP]
    return set(ws), {f"{a} {b}" for a, b in zip(ws, ws[1:])}


def pct(a, b):
    return round(a / b * 100, 1) if b else None


def keyword_sheet(rows):
    out = []
    by = defaultdict(list)
    for r in rows:
        by[(r["keyword"], r["platform"])].append(r)
    for (kw, plat), v in sorted(by.items(), key=lambda x: (x[0][0] or "", x[0][1])):
        c = Counter(r["_label"] for r in v)
        pos, neg = c["확정"] + c["상세 확인"], c["제외"]
        judged = pos + neg
        rate = pct(pos, judged)
        if judged < MIN_JUDGED:
            verdict = "판정 부족 — 더 쌓인 뒤 판단"
        elif pos == 0:
            verdict = "이 플랫폼에서 끄기 검토 (확정 0)"
        elif rate < 5:
            verdict = "좁히기 필요 — 제외어 후보 참고"
        elif rate >= 30:
            verdict = "좋음 — 페이지 늘리기 검토"
        else:
            verdict = "유지"
        out.append({"keyword": kw, "platform": plat, "model_type": v[0].get("model_type"), "n": len(v),
                    "sent": c["확정"], "review": c["상세 확인"], "excl": neg, "pending": c[""],
                    "rate": rate, "verdict": verdict})
    return out


def term_sheets(rows, cfg_keywords, common_exclude):
    """Model Type 별 제외어 후보 / 새 키워드 후보."""
    kw_by_type = defaultdict(list)
    excl_by_type = defaultdict(set)
    for k in cfg_keywords:
        kw_by_type[k.get("model_type")].append(k["keyword"].lower())
        excl_by_type[k.get("model_type")].update(w.lower() for w in k.get("exclude", []))
    common = {w.lower() for w in common_exclude}
    excl_out, new_out = [], []
    groups = defaultdict(list)
    for r in rows:
        if r["_label"]:
            groups[r.get("model_type") or "-"].append(r)
    for mt, v in sorted(groups.items()):
        pos_df, neg_df, ex = Counter(), Counter(), defaultdict(list)
        n_pos = sum(1 for r in v if r["_label"] != "제외")
        n_neg = len(v) - n_pos
        for r in v:
            one, two = tokens(r["title"])
            for t in one | two:
                (neg_df if r["_label"] == "제외" else pos_df)[t] += 1
                if len(ex[(t, r["_label"] == "제외")]) < 3:
                    ex[(t, r["_label"] == "제외")].append(r["title"])
        kws = " ".join(kw_by_type.get(mt, []))
        have_excl = excl_by_type.get(mt, set()) | common
        base = pct(n_pos, len(v)) or 0
        part = []
        for t, n in neg_df.items():  # 제외 상품명에 많고 확정·상세엔 거의 없음 (이미 제외어인 건 뺌)
            p = pos_df.get(t, 0)
            if n >= TERM_MIN_NEG and p <= max(1, n // 20) and " " not in t and t not in kws and t not in have_excl:
                part.append({"model_type": mt, "term": t, "neg": n, "pos": p, "neg_share": pct(n, n_neg),
                             "examples": " / ".join(ex[(t, True)])})
        excl_out += sorted(part, key=lambda x: -x["neg"])[:TOP_EXCL]
        part = []
        for t, p in pos_df.items():
            n = neg_df.get(t, 0)
            hit = pct(p, p + n)
            if p >= TERM_MIN_POS and hit >= 50 and hit >= base * TERM_LIFT and t not in kws:
                part.append({"model_type": mt, "term": t, "pos": p, "neg": n, "rate": hit, "base": base,
                             "pos_share": pct(p, n_pos), "examples": " / ".join(ex[(t, False)])})
        new_out += sorted(part, key=lambda x: -(x["pos"] * x["rate"]))[:TOP_NEW]
    return excl_out, new_out


def price_sheet(rows, cfg_keywords):
    caps = {}
    for k in cfg_keywords:
        caps.setdefault(k.get("model_type"), k.get("max_price"))
    by = defaultdict(list)
    for r in rows:
        if r["_label"] in ("확정", "상세 확인") and r.get("price"):
            by[r.get("model_type") or "-"].append(r["price"])
    out = []
    for mt, ps in sorted(by.items()):
        ps.sort()
        p90 = ps[min(len(ps) - 1, int(len(ps) * 0.9))]
        cap = caps.get(mt)
        over = sum(1 for p in ps if cap and p > cap)
        out.append({"model_type": mt, "n": len(ps), "min": ps[0], "median": int(median(ps)), "p90": p90,
                    "max": ps[-1], "cap": cap, "over": over,
                    "note": (f"확정 {over}건이 상한 초과 → 상한 올리기 검토" if over else "")})
    return out


def main():
    ap = argparse.ArgumentParser(description="검색 키워드 점검 리포트")
    ap.add_argument("--config", default=str(HERE / "keywords.local.json"))
    ap.add_argument("--since", help="이 날(YYYY-MM-DD) 이후 처음 본 상품만")
    ap.add_argument("--out", default=str(HERE / "output"))
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    cfg_keywords = [k for k in cfg["keywords"] if k.get("enabled", True)]

    rows = fetch(service_key(), args.since)
    total = len(rows)
    rows = [r for r in rows if not (r.get("reasons") or "").startswith(STORE_VIA)]  # 검색 결과만
    for r in rows:
        r["_label"] = label(r)
    c = Counter(r["_label"] for r in rows)
    print(f"허브 상품 {total:,}건 (매장 순회 제외 {len(rows):,}) — 확정 {c['확정']} / 상세 확인 {c['상세 확인']} / "
          f"제외 {c['제외']} / 미판정·자동 제외 {c['']}")

    ks = keyword_sheet(rows)
    excl, new = term_sheets(rows, cfg_keywords, cfg.get("common_exclude", []))
    prices = price_sheet(rows, cfg_keywords)
    picked = sorted((r for r in rows if r["_label"] in ("확정", "상세 확인")),
                    key=lambda r: (r.get("model_type") or "", r["keyword"] or "", r["_label"]))

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb.create_sheet("키워드 성적표"), [
        ("키워드", "keyword", 24), ("플랫폼", "platform", 10), ("Model Type", "model_type", 22), ("수집", "n", 8),
        ("확정", "sent", 7), ("상세 확인", "review", 9), ("제외", "excl", 7), ("미판정", "pending", 8),
        ("적중률(%)", "rate", 10), ("판단", "verdict", 34)], ks)
    _sheet(wb.create_sheet("제외어 후보"), [
        ("Model Type", "model_type", 22), ("단어", "term", 14), ("제외 상품명", "neg", 10), ("확정·상세", "pos", 9),
        ("제외 중 비율(%)", "neg_share", 13), ("예시 (제외된 상품명)", "examples", 90)], excl)
    _sheet(wb.create_sheet("새 키워드 후보"), [
        ("Model Type", "model_type", 22), ("단어·묶음", "term", 18), ("확정·상세", "pos", 9), ("제외", "neg", 7),
        ("적중률(%)", "rate", 10), ("유형 평균(%)", "base", 11), ("확정 중 비율(%)", "pos_share", 13),
        ("예시 (확정·상세 상품명)", "examples", 90)], new)
    _sheet(wb.create_sheet("가격 상한"), [
        ("Model Type", "model_type", 22), ("확정·상세", "n", 9), ("최저", "min", 9), ("중간값", "median", 9),
        ("상위 10%", "p90", 9), ("최고", "max", 9), ("설정 상한", "cap", 10), ("상한 초과", "over", 9),
        ("메모", "note", 34)], prices)
    _sheet(wb.create_sheet("확정·상세 상품명"), [
        ("판정", "_label", 9), ("Model Type", "model_type", 22), ("키워드", "keyword", 22), ("플랫폼", "platform", 10),
        ("상품명", "title", 70), ("가격(원)", "price", 10), ("판매자", "seller", 16), ("상품 URL", "url", 50)], picked)
    _sheet(wb.create_sheet("현재 키워드 설정"), [
        ("Model Type", "model_type", 22), ("키워드", "keyword", 24), ("가격 상한", "max_price", 10),
        ("제외어", "exclude_s", 50)], [{**k, "exclude_s": ", ".join(k.get("exclude", []))} for k in cfg_keywords])

    out = Path(args.out) / f"keyword_review_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    flags = Counter(k["verdict"].split(" ")[0] for k in ks)
    print(f"키워드×플랫폼 {len(ks)}개: " + ", ".join(f"{k} {v}" for k, v in flags.most_common()))
    print(f"제외어 후보 {len(excl)}개 / 새 키워드 후보 {len(new)}개")
    print(f"저장: {out}")


if __name__ == "__main__":
    main()
