import os
import threading
import queue
import traceback
import time
from dataclasses import dataclass
from typing import Any, List, Optional, Tuple

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait, Select as SeleniumSelect
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, StaleElementReferenceException


def _detect_chrome_major():
    """Registry 에서 설치된 Chrome major 버전 읽기. UC version_main=None 이 CfT 최신 드라이버를 받아오는 동작 회피용. 실패 시 None → UC 기본 동작."""
    try:
        import winreg
        for hive, path in (
            (winreg.HKEY_CURRENT_USER, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Wow6432Node\Google\Chrome\BLBeacon"),
        ):
            try:
                with winreg.OpenKey(hive, path) as key:
                    version, _ = winreg.QueryValueEx(key, "version")
                    return int(str(version).split(".")[0])
            except OSError:
                continue
    except Exception:
        pass
    return None


APP_TITLE = "Gmarket VeRO 신고 매크로"
VERO_URL  = "https://www.gmarket.co.kr/gmap/vero_main.asp"

YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF59D")
GREEN_FILL  = PatternFill(fill_type="solid", fgColor="C8E6C9")
RED_FILL    = PatternFill(fill_type="solid", fgColor="FFCDD2")

REPORT_TYPE_MAP = {
    "위조품": "1", "위조상품": "1",
    "저작권 침해": "2", "저작권침해": "2",
    "상표권 침해": "3", "상표권침해": "3",
    "특허.실용신안권침해": "4", "특허·실용신안권침해": "4",
    "특허/실용신안권침해": "4", "특허권침해": "4", "실용신안권침해": "4",
    "디자인권 침해": "5", "디자인권침해": "5",
    "초상및성명권침해": "6", "초상 및 성명권침해": "6",
    "기타": "7",
}


# ─── 데이터 클래스 ─────────────────────────────────────────────────────────────

@dataclass
class ReportItem:
    row_no: int
    product_no: str
    report_type_text: str
    product_url: str
    status: str = "대기"
    detail: str = ""
    result_fill: Optional[str] = None


# ─── 셀렉터 정의 ──────────────────────────────────────────────────────────────

class Selectors:
    LOGIN_ID = [
        (By.CSS_SELECTOR, 'input#id'),
        (By.CSS_SELECTOR, 'input[name="id"]'),
    ]
    LOGIN_PW = [
        (By.CSS_SELECTOR, 'input#pwd'),
        (By.CSS_SELECTOR, 'input[name="pwd"]'),
        (By.CSS_SELECTOR, 'input[type="password"]'),
    ]
    LOGIN_SUBMIT = [
        (By.CSS_SELECTOR, 'a[href="javascript:go();"]'),
        (By.XPATH, '//a[.//img[contains(@src,"btn_login.gif")]]'),
    ]
    LOGIN_SUCCESS = [
        (By.CSS_SELECTOR, 'a[href="javascript:logout();"]'),
        (By.CSS_SELECTOR, 'img[src*="btn_logout.gif"]'),
        (By.XPATH, '//*[contains(text(),"VeRO회원")]'),
        (By.XPATH, '//*[contains(text(),"환영합니다")]'),
    ]
    SEARCH_KEYWORD         = (By.CSS_SELECTOR, 'input#search_str')
    SEARCH_BUTTON          = [
        (By.CSS_SELECTOR, 'a[href="javascript:go(\'상품검색\',1);"]'),
        (By.XPATH, '//a[@href="javascript:go(\'상품검색\',1);"]'),
        (By.XPATH, '//input[@id="search_str"]/following-sibling::a[1]'),
        (By.XPATH, '//td[.//input[@id="search_str"]]//a[@href="javascript:go(\'상품검색\',1);"]'),
        (By.XPATH, '//a[.//img[contains(@src,"btn_s01")]]'),
        (By.XPATH, '//a[contains(@href,"상품검색")]'),
    ]
    SEARCH_RESULT_CHECKBOX = (By.CSS_SELECTOR, 'input[name="singo"]')
    REPORT_TYPE_SELECT     = (By.CSS_SELECTOR, 'select#report_type')
    REPORT_SUBMIT          = [
        (By.CSS_SELECTOR, 'a[href*="신고하기"]'),
        (By.XPATH, '//a[contains(@href,"신고하기")]'),
    ]
    PRODUCT_NO_CHECKBOX    = [
        (By.CSS_SELECTOR, 'input#cbx_seller_nm'),
    ]


# ─── 유틸 ──────────────────────────────────────────────────────────────────────

def normalize_text(text: str) -> str:
    return (text or "").strip().replace(" ", "").replace("\u00a0", "")


def get_report_type_value(label: str) -> Optional[str]:
    n = normalize_text(label)
    for k, v in REPORT_TYPE_MAP.items():
        if normalize_text(k) == n:
            return v
    return None


def safe_str(value) -> str:
    return "" if value is None else str(value).strip()


# ─── Selenium 프레임 탐색 헬퍼 ────────────────────────────────────────────────

def reset_frame(driver: uc.Chrome):
    try:
        driver.switch_to.default_content()
    except Exception:
        pass


def get_iframe_srcs(driver: uc.Chrome) -> List[Tuple[int, str]]:
    reset_frame(driver)
    try:
        iframes = driver.find_elements(By.TAG_NAME, 'iframe')
        return [(i, el.get_attribute('src') or '') for i, el in enumerate(iframes)]
    except Exception:
        return []


def _switch_to_frame_idx(driver: uc.Chrome, idx: int):
    reset_frame(driver)
    if idx >= 0:
        iframes = driver.find_elements(By.TAG_NAME, 'iframe')
        if idx < len(iframes):
            driver.switch_to.frame(iframes[idx])


def _quick_find(driver: uc.Chrome, by: str, selector: str, timeout: float = 0.5) -> Any:
    """현재 프레임 컨텍스트에서 빠르게 요소 탐색. 못 찾으면 None."""
    try:
        return WebDriverWait(driver, timeout).until(
            EC.visibility_of_element_located((by, selector))
        )
    except Exception:
        return None


def _quick_find_present(driver: uc.Chrome, by: str, selector: str, timeout: float = 0.5) -> Any:
    """현재 프레임 컨텍스트에서 DOM 존재만 확인 (visibility 불필요)."""
    try:
        return WebDriverWait(driver, timeout).until(
            EC.presence_of_element_located((by, selector))
        )
    except Exception:
        return None


def find_element_in_frames(
    driver: uc.Chrome,
    by: str,
    selector: str,
    visible: bool = True,
    timeout: float = 0.8,
) -> Tuple[Optional[int], Any]:
    cond = EC.visibility_of_element_located if visible else EC.presence_of_element_located

    def _try(idx: int):
        try:
            el = WebDriverWait(driver, timeout).until(cond((by, selector)))
            return idx, el
        except Exception:
            return None, None

    reset_frame(driver)
    fi, el = _try(-1)
    if el is not None:
        return fi, el

    reset_frame(driver)
    iframes = driver.find_elements(By.TAG_NAME, 'iframe')
    for i, iframe in enumerate(iframes):
        reset_frame(driver)
        try:
            driver.switch_to.frame(iframe)
        except Exception:
            continue
        fi, el = _try(i)
        if el is not None:
            return fi, el

    reset_frame(driver)
    return None, None


def find_first_in_frames(
    driver: uc.Chrome,
    candidates: List[Tuple[str, str]],
    visible: bool = True,
    timeout: float = 0.8,
) -> Tuple[Optional[int], Any, Optional[Tuple]]:
    for by, sel in candidates:
        fi, el = find_element_in_frames(driver, by, sel, visible=visible, timeout=timeout)
        if el is not None:
            return fi, el, (by, sel)
    return None, None, None


def accept_alert(driver: uc.Chrome, timeout: float = 3.0) -> Optional[str]:
    try:
        alert = WebDriverWait(driver, timeout).until(EC.alert_is_present())
        msg = alert.text
        alert.accept()
        return msg
    except Exception:
        return None


def safe_fill(element, text: str):
    element.clear()
    element.send_keys(text)


def ensure_checked(driver: uc.Chrome, element):
    try:
        if not element.is_selected():
            element.click()
    except Exception:
        try:
            driver.execute_script("arguments[0].click();", element)
        except Exception:
            driver.execute_script(
                "arguments[0].checked=true;"
                "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                element,
            )


# ─── VeRO 신고 엔진 ────────────────────────────────────────────────────────────

class VeroReporter:
    def __init__(self, user_id: str, password: str, log_callback=None):
        self.user_id  = user_id
        self.password = password
        self.log_callback = log_callback or (lambda m: None)
        self.driver: Optional[uc.Chrome] = None
        # 신고 페이지가 들어있는 iframe 인덱스 캐시 (-1=기본콘텐츠, None=미확인)
        self._report_frame_idx: Optional[int] = None

    def log(self, msg: str):
        self.log_callback(msg)

    # ── 캐시된 프레임으로 빠르게 전환 ──────────────────────────────────────────

    def _go_report_frame(self):
        """캐시된 신고 페이지 프레임으로 전환."""
        if self._report_frame_idx is not None:
            _switch_to_frame_idx(self.driver, self._report_frame_idx)

    def _fast_find(self, by: str, selector: str, timeout: float = 0.5) -> Any:
        """캐시된 프레임에서 빠르게 요소 탐색."""
        self._go_report_frame()
        return _quick_find(self.driver, by, selector, timeout)

    def _fast_find_candidates(self, candidates: List[Tuple[str, str]],
                              timeout: float = 0.5) -> Any:
        """캐시된 프레임에서 복수 셀렉터 중 첫 매칭 요소 반환."""
        self._go_report_frame()
        for by, sel in candidates:
            el = _quick_find(self.driver, by, sel, timeout)
            if el is not None:
                return el
        return None

    # ── 드라이버 생명주기 ───────────────────────────────────────────────────────

    def start(self):
        opts = uc.ChromeOptions()
        opts.add_argument('--window-size=1440,1200')
        opts.add_argument('--no-sandbox')
        opts.add_argument('--disable-dev-shm-usage')
        self.driver = uc.Chrome(
            headless=False,
            options=opts,
            use_subprocess=True,  # frozen launcher 환경에서 통합 런처 재실행 방지
            version_main=_detect_chrome_major(),  # CfT 최신 드라이버 자동 다운로드 회피
        )
        self.driver.implicitly_wait(0)
        self.driver.set_page_load_timeout(30)

    def close(self):
        try:
            if self.driver:
                self.driver.quit()
        except Exception:
            pass

    # ── 로그인 ─────────────────────────────────────────────────────────────────

    def login(self):
        self.log("[접속] VeRO 페이지 접속")
        self.driver.get(VERO_URL)
        time.sleep(2)

        for i, src in get_iframe_srcs(self.driver):
            self.log(f"[프레임] {i}: {src}")

        fi_id, el_id, sel_id = find_first_in_frames(
            self.driver, Selectors.LOGIN_ID, timeout=3.0
        )
        fi_pw, el_pw, sel_pw = find_first_in_frames(
            self.driver, Selectors.LOGIN_PW, timeout=3.0
        )

        if el_id is None:
            raise RuntimeError("ID 입력칸을 찾지 못했습니다.")
        if el_pw is None:
            raise RuntimeError("PW 입력칸을 찾지 못했습니다.")

        self.log(f"[로그인] ID selector={sel_id}  frame_idx={fi_id}")
        self.log(f"[로그인] PW selector={sel_pw}  frame_idx={fi_pw}")

        _switch_to_frame_idx(self.driver, fi_id)
        _, el_id2, _ = find_first_in_frames(self.driver, Selectors.LOGIN_ID, timeout=2)
        if el_id2:
            safe_fill(el_id2, self.user_id)

        _switch_to_frame_idx(self.driver, fi_pw)
        fi_pw2, el_pw2, _ = find_first_in_frames(self.driver, Selectors.LOGIN_PW, timeout=2)
        if el_pw2:
            safe_fill(el_pw2, self.password)

        submit_errors: List[str] = []
        submitted = False

        fi_sub, el_sub, sel_sub = find_first_in_frames(
            self.driver, Selectors.LOGIN_SUBMIT, timeout=3.0
        )
        self.log(f"[로그인] BTN selector={sel_sub}  frame_idx={fi_sub}")

        if el_sub:
            try:
                el_sub.click()
                self.log("[로그인] 버튼 클릭 완료")
                submitted = True
            except Exception as e:
                submit_errors.append(f"click 실패: {e}")

        if not submitted:
            try:
                _switch_to_frame_idx(self.driver, fi_id if fi_id is not None else -1)
                self.driver.execute_script("go()")
                self.log("[로그인] go() JS 실행 완료")
                submitted = True
            except Exception as e:
                submit_errors.append(f"go() 실패: {e}")

        if not submitted:
            try:
                _switch_to_frame_idx(self.driver, fi_pw2 if fi_pw2 is not None else -1)
                _, el_pw3, _ = find_first_in_frames(self.driver, Selectors.LOGIN_PW, timeout=2)
                if el_pw3:
                    el_pw3.send_keys(Keys.RETURN)
                    self.log("[로그인] Enter 제출 완료")
                    submitted = True
            except Exception as e:
                submit_errors.append(f"Enter 실패: {e}")

        accept_alert(self.driver, timeout=2)
        time.sleep(3.5)

        success_used = None

        fi_s, el_s, sel_s = find_first_in_frames(
            self.driver, Selectors.LOGIN_SUCCESS, visible=True, timeout=2
        )
        if el_s:
            success_used = f"{sel_s} (visible, frame_idx={fi_s})"

        if not success_used:
            fi_s, el_s, sel_s = find_first_in_frames(
                self.driver, Selectors.LOGIN_SUCCESS, visible=False, timeout=1
            )
            if el_s:
                success_used = f"{sel_s} (attached, frame_idx={fi_s})"

        if not success_used:
            reset_frame(self.driver)
            iframes = self.driver.find_elements(By.TAG_NAME, 'iframe')
            for iframe in iframes:
                src = iframe.get_attribute('src') or ''
                if 'vero_login_check.asp' in src:
                    try:
                        self.driver.switch_to.frame(iframe)
                        body_text = self.driver.find_element(By.TAG_NAME, 'body').text
                        if any(kw in body_text for kw in ['VeRO회원', '환영합니다', 'logout']):
                            success_used = "vero_login_check.asp 성공 문구"
                    except Exception:
                        pass
                    finally:
                        reset_frame(self.driver)
                    break

        if not success_used:
            for i, src in get_iframe_srcs(self.driver):
                self.log(f"[로그인후 프레임] {i}: {src}")
            joined = " | ".join(submit_errors) if submit_errors else "제출 시도 기록 없음"
            time.sleep(20)
            raise RuntimeError(
                f"로그인 성공을 확인하지 못했습니다. 아이디·비밀번호를 다시 확인하세요. "
                f"제출 로그: {joined}"
            )

        self.log(f"[로그인] 성공 확인: {success_used}")

    # ── 신고 페이지 이동 ────────────────────────────────────────────────────────

    def goto_report_page(self):
        self.log("[이동] 신고 페이지 이동")

        reset_frame(self.driver)
        iframes = self.driver.find_elements(By.TAG_NAME, 'iframe')
        left_idx = next(
            (i for i, el in enumerate(iframes)
             if 'vero_left.asp' in (el.get_attribute('src') or '')),
            None,
        )

        if left_idx is not None:
            self.driver.switch_to.frame(iframes[left_idx])
            try:
                self.driver.execute_script("CallVeroLayerFile('vero_report')")
                self.log("[이동] left frame 에서 CallVeroLayerFile 실행")
            except Exception:
                reset_frame(self.driver)
                self.driver.execute_script("CallVeroLayerFile('vero_report')")
                self.log("[이동] top page 에서 CallVeroLayerFile 실행 (fallback)")
        else:
            reset_frame(self.driver)
            self.driver.execute_script("CallVeroLayerFile('vero_report')")
            self.log("[이동] top page 에서 CallVeroLayerFile 실행")

        time.sleep(2)

        # 신고 페이지 프레임 탐색 & 캐시
        by, sel = Selectors.SEARCH_KEYWORD
        fi, el = find_element_in_frames(self.driver, by, sel, visible=True, timeout=2)
        if el is None:
            for i, src in get_iframe_srcs(self.driver):
                self.log(f"[신고페이지 프레임] {i}: {src}")
            raise RuntimeError("신고 페이지 검색 입력칸을 찾지 못했습니다.")

        # ★ 프레임 인덱스 캐시 — 이후 모든 신고 동작에서 재사용
        self._report_frame_idx = fi
        self.log(f"[이동] 신고 페이지 진입 완료 / 캐시 frame_idx={fi}")

    # ── 상품 검색 ──────────────────────────────────────────────────────────────

    def search_product(self, product_no: str):
        # 상품번호 라디오/체크박스 선택
        el_cb = self._fast_find_candidates(Selectors.PRODUCT_NO_CHECKBOX, timeout=0.8)
        if el_cb is None:
            raise RuntimeError("상품번호 체크박스를 찾지 못했습니다.")
        ensure_checked(self.driver, el_cb)

        # 키워드 입력
        by_kw, sel_kw = Selectors.SEARCH_KEYWORD
        el_kw = self._fast_find(by_kw, sel_kw, timeout=0.8)
        if el_kw is None:
            raise RuntimeError("검색 입력칸을 찾지 못했습니다.")

        self._go_report_frame()
        try:
            el_kw.click()
        except Exception:
            pass
        try:
            el_kw.clear()
        except Exception:
            self.driver.execute_script("arguments[0].value = '';", el_kw)
        el_kw.send_keys(str(product_no))
        self.log(f"[검색] 상품번호 입력: {product_no}")

        searched = False
        errors = []

        # 1) 입력창 기준으로 가장 가까운 a 버튼을 찾아 실제 클릭
        try:
            self._go_report_frame()
            btn = self.driver.find_element(
                By.XPATH,
                '//input[@id="search_str"]/following-sibling::a[1]'
            )
            self.driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
            try:
                btn.click()
            except Exception:
                self.driver.execute_script("arguments[0].click();", btn)
            self.log("[검색] input 옆 검색 버튼 클릭 완료")
            searched = True
        except Exception as e:
            errors.append(f"neighbor_button_click 실패: {e}")
            self.log(f"[검색] neighbor_button_click 실패: {e}")

        # 2) href가 정확히 일치하는 버튼 클릭
        if not searched:
            try:
                self._go_report_frame()
                btn = self.driver.find_element(
                    By.XPATH,
                    '//a[@href="javascript:go(\'상품검색\',1);"]'
                )
                self.driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
                try:
                    btn.click()
                except Exception:
                    self.driver.execute_script("arguments[0].click();", btn)
                self.log("[검색] href 일치 검색 버튼 클릭 완료")
                searched = True
            except Exception as e:
                errors.append(f"exact_href_click 실패: {e}")
                self.log(f"[검색] exact_href_click 실패: {e}")

        # 3) onclick 대신 페이지 함수 직접 호출
        if not searched:
            try:
                self._go_report_frame()
                self.driver.execute_script("go('상품검색',1);")
                self.log("[검색] JS go('상품검색',1) 실행 완료")
                searched = True
            except Exception as e:
                errors.append(f"direct_go_call 실패: {e}")
                self.log(f"[검색] direct_go_call 실패: {e}")

        # 4) 링크의 href 내용을 그대로 eval 실행
        if not searched:
            try:
                self._go_report_frame()
                ok = self.driver.execute_script("""
                    var a = document.evaluate(
                        "//input[@id='search_str']/following-sibling::a[1]",
                        document,
                        null,
                        XPathResult.FIRST_ORDERED_NODE_TYPE,
                        null
                    ).singleNodeValue;
                    if (!a) {
                        a = document.querySelector("a[href=\"javascript:go('상품검색',1);\"]");
                    }
                    if (!a) return false;
                    var href = a.getAttribute('href') || '';
                    if (href.toLowerCase().startsWith('javascript:')) {
                        eval(href.substring('javascript:'.length));
                        return true;
                    }
                    return false;
                """)
                if ok:
                    self.log("[검색] href eval 실행 완료")
                    searched = True
                else:
                    errors.append("href_eval 실패: 버튼 미발견")
            except Exception as e:
                errors.append(f"href_eval 실패: {e}")
                self.log(f"[검색] href_eval 실패: {e}")

        if not searched:
            raise RuntimeError("검색 버튼 실행에 실패했습니다. / " + " | ".join(errors))

        # 결과 갱신 대기
        try:
            self._go_report_frame()
            WebDriverWait(self.driver, 3).until(
                lambda d: len(d.find_elements(*Selectors.SEARCH_RESULT_CHECKBOX)) > 0
                or (self.driver.find_element(*Selectors.SEARCH_KEYWORD).get_attribute("value") or "").strip() == str(product_no).strip()
            )
        except Exception:
            pass

        time.sleep(1.5)

    def has_search_result(self, product_no: str = "") -> bool:
        try:
            by, sel = Selectors.SEARCH_RESULT_CHECKBOX
            el = self._fast_find(by, sel, timeout=1.0)
            if el is None:
                return False
            if not product_no:
                return True
            val = (el.get_attribute("value") or "").split("|")[0].strip()
            return val == str(product_no).strip()
        except Exception:
            return False

    def select_first_result(self):
        by, sel = Selectors.SEARCH_RESULT_CHECKBOX
        el = self._fast_find(by, sel, timeout=1.0)
        if el is None:
            raise RuntimeError("검색 결과의 첫 행 체크박스를 찾지 못했습니다.")
        ensure_checked(self.driver, el)
        self.log("[검색결과] 첫 행 선택 체크 완료")

    def select_report_type(self, report_type_text: str):
        value = get_report_type_value(report_type_text)
        if not value:
            supported = ", ".join(sorted(REPORT_TYPE_MAP.keys()))
            raise RuntimeError(
                f"지원하지 않는 침해유형: '{report_type_text}'\n"
                f"지원값: {supported}\n엑셀 B열 값을 확인해주세요."
            )
        by, sel = Selectors.REPORT_TYPE_SELECT
        el = self._fast_find(by, sel, timeout=0.5)
        if el is None:
            raise RuntimeError("침해유형 드롭다운을 찾지 못했습니다.")
        SeleniumSelect(el).select_by_value(value)
        self.log(f"[신고] 침해유형 선택: {report_type_text} ({value})")

    def submit_report(self):
        # ★ JS 직접 호출 우선
        submitted = False
        try:
            self._go_report_frame()
            self.driver.execute_script("go('신고하기')")
            self.log("[신고] JS go('신고하기') 실행 완료")
            submitted = True
        except Exception as e:
            self.log(f"[신고] JS 실행 실패: {e}")

        # 버튼 클릭 폴백
        if not submitted:
            el = self._fast_find_candidates(Selectors.REPORT_SUBMIT, timeout=0.8)
            if el is None:
                raise RuntimeError("선택신고하기 버튼을 찾지 못했습니다.")
            try:
                el.click()
            except Exception:
                self.driver.execute_script("arguments[0].click();", el)
            self.log("[신고] 버튼 클릭 완료")

        accept_alert(self.driver, timeout=5)
        time.sleep(1.0)
        self.log("[신고] 신고 처리 완료")

    def report_one(self, item: ReportItem) -> Tuple[str, str]:
        try:
            self.search_product(item.product_no)
            if not self.has_search_result(item.product_no):
                return "진행 불가", "검색 결과 없음 또는 판매 중이 아닌 게시물"
            self.select_first_result()
            self.select_report_type(item.report_type_text)
            self.submit_report()
            return "신고 완료", "정상 처리됨"
        except TimeoutException as e:
            return "실패", f"타임아웃: {e}"
        except Exception as e:
            return "실패", str(e)


# ─── UI ────────────────────────────────────────────────────────────────────────

class LoginDialog(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("VeRO 로그인 정보 입력")
        self.resizable(False, False)
        self.result = None
        self.transient(parent)
        self.grab_set()

        frame = ttk.Frame(self, padding=12)
        frame.pack(fill="both", expand=True)

        ttk.Label(frame, text="ID").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Label(frame, text="PW").grid(row=1, column=0, sticky="w", pady=4)

        self.id_var = tk.StringVar()
        self.pw_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.id_var, width=30).grid(row=0, column=1, pady=4)
        ttk.Entry(frame, textvariable=self.pw_var, width=30, show="*").grid(row=1, column=1, pady=4)

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=2, column=0, columnspan=2, pady=(12, 0))
        ttk.Button(btn_frame, text="확인", command=self.on_ok).pack(side="left", padx=4)
        ttk.Button(btn_frame, text="취소", command=self.on_cancel).pack(side="left", padx=4)

        self.bind("<Return>", lambda e: self.on_ok())
        self.bind("<Escape>", lambda e: self.on_cancel())
        self.update_idletasks()
        self.update_idletasks()
        w = self.winfo_width()
        h = self.winfo_height()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = int((sw / 2) - (w / 2))
        y = int((sh / 2) - (h / 2))
        self.geometry(f"+{x}+{y}")

    def on_ok(self):
        uid = self.id_var.get().strip()
        pw  = self.pw_var.get().strip()
        if not uid or not pw:
            messagebox.showwarning("입력 필요", "ID와 PW를 모두 입력해주세요.", parent=self)
            return
        self.result = (uid, pw)
        self.destroy()

    def on_cancel(self):
        self.result = None
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("980x620")
        self.center_window(980, 620)

        self.items: List[ReportItem] = []
        self.excel_path: Optional[str] = None
        self.log_queue = queue.Queue()
        self.worker_thread: Optional[threading.Thread] = None
        self.progress_var = tk.DoubleVar(value=0)
        self.progress_text_var = tk.StringVar(value="진행률: 0 / 0")

        self._build_ui()
        self.after(150, self._poll_log_queue)


    def center_window(self, width, height):
        self.update_idletasks()
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        x = int((screen_width / 2) - (width / 2))
        y = int((screen_height / 2) - (height / 2))
        self.geometry(f"{width}x{height}+{x}+{y}")
    def _build_ui(self):
        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")

        ttk.Button(top, text="엑셀 업로드",   command=self.load_excel).pack(side="left", padx=4)
        ttk.Button(top, text="신고 진행",      command=self.start_report).pack(side="left", padx=4)
        ttk.Button(top, text="결과 엑셀 저장", command=self.export_result).pack(side="left", padx=4)
        self.file_label = ttk.Label(top, text="업로드된 파일 없음")
        self.file_label.pack(side="left", padx=10)

        progress_frame = ttk.Frame(self, padding=(10, 0, 10, 8))
        progress_frame.pack(fill="x")
        self.progress_bar = ttk.Progressbar(
            progress_frame,
            variable=self.progress_var,
            maximum=100,
            mode="determinate",
        )
        self.progress_bar.pack(fill="x", side="left", expand=True)
        ttk.Label(progress_frame, textvariable=self.progress_text_var, width=18).pack(side="left", padx=(10, 0))

        mid = ttk.Frame(self, padding=(10, 0, 10, 0))
        mid.pack(fill="both", expand=True)

        cols = ("row_no", "product_no", "report_type", "product_url", "status", "detail")
        self.tree = ttk.Treeview(mid, columns=cols, show="headings", height=18)
        hdrs  = dict(row_no="엑셀행", product_no="상품번호", report_type="침해유형",
                     product_url="상품 URL", status="상태", detail="상세결과")
        widths = dict(row_no=70, product_no=120, report_type=130,
                      product_url=260, status=100, detail=250)
        for c in cols:
            self.tree.heading(c, text=hdrs[c])
            self.tree.column(c, width=widths[c], anchor="w")

        vsb = ttk.Scrollbar(mid, orient="vertical",   command=self.tree.yview)
        hsb = ttk.Scrollbar(mid, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        mid.grid_rowconfigure(0, weight=1)
        mid.grid_columnconfigure(0, weight=1)

        bottom = ttk.LabelFrame(self, text="로그", padding=10)
        bottom.pack(fill="both", expand=False, padx=10, pady=10)
        self.log_text = tk.Text(bottom, height=12, wrap="word")
        self.log_text.pack(fill="both", expand=True)

    def log(self, msg: str):
        self.log_queue.put(msg)

    def _poll_log_queue(self):
        while not self.log_queue.empty():
            msg = self.log_queue.get_nowait()
            self.log_text.insert("end", msg + "\n")
            self.log_text.see("end")
        self.after(150, self._poll_log_queue)


    def reset_progress(self):
        total = len(self.items)
        self.progress_var.set(0)
        self.progress_text_var.set(f"진행률: 0 / {total}")

    def update_progress(self, completed: int, total: int):
        percent = (completed / total * 100) if total else 0
        self.progress_var.set(percent)
        self.progress_text_var.set(f"진행률: {completed} / {total}")

    def refresh_tree(self):
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        for item in self.items:
            self.tree.insert("", "end", values=(
                item.row_no, item.product_no, item.report_type_text,
                item.product_url, item.status, item.detail,
            ))

    def load_excel(self):
        path = filedialog.askopenfilename(
            title="엑셀 템플릿 선택",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if not path:
            return
        try:
            wb = load_workbook(path)
            ws = wb.active
            items: List[ReportItem] = []
            row = 2
            while True:
                pno  = safe_str(ws[f"A{row}"].value)
                rtyp = safe_str(ws[f"B{row}"].value)
                purl = safe_str(ws[f"D{row}"].value)
                if not pno and not rtyp and not purl:
                    break
                if pno:
                    items.append(ReportItem(row_no=row, product_no=pno,
                                            report_type_text=rtyp, product_url=purl))
                row += 1
            self.items = items
            self.excel_path = path
            self.file_label.config(text=os.path.basename(path))
            self.refresh_tree()
            self.reset_progress()
            self.log(f"[엑셀] {len(items)}건 로드 완료")
            messagebox.showinfo("완료", f"{len(items)}건을 불러왔습니다.")
        except Exception as e:
            messagebox.showerror("오류", f"엑셀 로드 실패\n{e}")

    def update_item_result(self, index: int, status: str, detail: str):
        self.items[index].status = status
        self.items[index].detail = detail
        self.items[index].result_fill = (
            "green" if status == "신고 완료" else
            "yellow" if status == "진행 불가" else "red"
        )
        self.after(0, self.refresh_tree)

    def start_report(self):
        if not self.items:
            messagebox.showwarning("대상 없음", "먼저 엑셀 파일을 업로드해주세요.")
            return
        if self.worker_thread and self.worker_thread.is_alive():
            messagebox.showwarning("진행 중", "이미 작업이 진행 중입니다.")
            return

        dlg = LoginDialog(self)
        self.wait_window(dlg)
        if not dlg.result:
            return
        user_id, pw = dlg.result

        for item in self.items:
            item.status = "대기"
            item.detail = ""
            item.result_fill = None
        self.refresh_tree()
        self.reset_progress()

        self.worker_thread = threading.Thread(
            target=self._run_report_worker,
            args=(user_id, pw),
            daemon=True,
        )
        self.worker_thread.start()

    def _run_report_worker(self, user_id: str, pw: str):
        reporter = VeroReporter(user_id, pw, log_callback=self.log)
        success = skipped = failed = 0
        try:
            reporter.start()
            reporter.login()
            reporter.goto_report_page()

            total = len(self.items)
            for idx, item in enumerate(self.items):
                self.log(f"[처리] 시작: 행 {item.row_no} / 상품번호 {item.product_no}")
                status, detail = reporter.report_one(item)
                self.update_item_result(idx, status, detail)
                if status == "신고 완료":   success += 1
                elif status == "진행 불가": skipped += 1
                else:                       failed  += 1
                self.after(0, lambda completed=idx + 1, total=total: self.update_progress(completed, total))
                self.log(f"[처리] 결과: {item.product_no} → {status} / {detail}")

            summary = (f"전체 {len(self.items)}건\n"
                       f"신고 완료: {success}건\n진행 불가: {skipped}건\n실패: {failed}건")
            self.log("[완료] " + summary.replace("\n", " / "))
            self.after(0, lambda: messagebox.showinfo("신고 작업 완료", summary))
        except Exception as e:
            err_msg = str(e)
            self.log("[오류] " + err_msg)
            self.log(traceback.format_exc())
            self.after(0, lambda m=err_msg: messagebox.showerror("실행 오류", m))
        finally:
            reporter.close()

    def export_result(self):
        if not self.items:
            messagebox.showwarning("대상 없음", "먼저 엑셀을 업로드해주세요.")
            return
        if not self.excel_path:
            messagebox.showwarning("원본 없음", "원본 엑셀 경로를 찾을 수 없습니다.")
            return

        save_path = filedialog.asksaveasfilename(
            title="결과 엑셀 저장",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="vero_report_result.xlsx",
        )
        if not save_path:
            return

        try:
            wb = load_workbook(self.excel_path)
            ws = wb.active
            ws.cell(row=1, column=5, value="처리상태")
            ws.cell(row=1, column=6, value="상세결과")

            fill_map = {"yellow": YELLOW_FILL, "green": GREEN_FILL, "red": RED_FILL}
            for item in self.items:
                ws.cell(row=item.row_no, column=5, value=item.status)
                ws.cell(row=item.row_no, column=6, value=item.detail)
                if item.result_fill in fill_map:
                    ws[f"A{item.row_no}"].fill = fill_map[item.result_fill]

            wb.save(save_path)
            messagebox.showinfo("저장 완료", f"결과 파일 저장 완료\n{save_path}")
            self.log(f"[엑셀] 결과 저장 완료: {save_path}")
        except Exception as e:
            messagebox.showerror("저장 오류", f"결과 저장 실패\n{e}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
