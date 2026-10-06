import os
import time
import random
import traceback
import datetime
from io import BytesIO
from PIL import Image as PilImage
import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext


def _detect_chrome_major():
    """Registry 에서 설치된 Chrome major 버전 읽기. UC version_main=None 이 CfT 최신 드라이버를 받아오는 동작 회피용. 실패 시 None → UC 기본 동작."""
    try:
        import winreg
        for hive, path in (
            (winreg.HKEY_CURRENT_USER, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Wow6432Node\Google\Chrome\BLBeacon"),
        ):
            try:
                with winreg.OpenKey(hive, path) as key:
                    version, _ = winreg.QueryValueEx(key, "version")
                    return int(str(version).split(".")[0])
            except OSError:
                continue
    except Exception:
        pass
    return None


# 헬퍼(Helper) 함수
def find_and_click_show_more(driver):
    """(11번가 전용) '더보기' 버튼 클릭 함수"""
    try:
        wait = WebDriverWait(driver, 3)
        button = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[@id='prdDetailDescMore']")))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", button)
        time.sleep(0.5)
        driver.execute_script("arguments[0].click();", button)
        return True
    except Exception:
        return False

def start_hide_fixed_ui_elements_interval(driver, selectors_to_hide, interval=200):
    """캡처 시 UI 숨기는 JS 실행"""
    js_selectors_array = str(selectors_to_hide)
    js_code = f"""
    if (!window.__HIDE_UI_INTERVAL) {{
      window.__HIDE_UI_INTERVAL = setInterval(function() {{
        let selectors = {js_selectors_array};
        for(let s of selectors){{
          try {{
            let elems = document.querySelectorAll(s);
            elems.forEach(e => {{ e.style.display='none'; e.style.visibility='hidden'; }});
          }} catch (err) {{}}
        }}
      }}, {interval});
    }}
    """
    driver.execute_script(js_code)

def stop_hide_fixed_ui_elements_interval(driver):
    """UI 숨김 반복 중단"""
    js_code = "if (window.__HIDE_UI_INTERVAL) { clearInterval(window.__HIDE_UI_INTERVAL); window.__HIDE_UI_INTERVAL = null; }"
    driver.execute_script(js_code)

def hide_elements_by_xpath(driver, xpaths):
    """XPath 목록으로 UI 숨기기"""
    for xpath in xpaths:
        try:
            elements_to_hide = driver.find_elements(By.XPATH, xpath)
            for element in elements_to_hide:
                driver.execute_script("arguments[0].style.display = 'none';", element)
        except Exception:
            continue

def hide_scrollbar(driver):
    driver.execute_script("document.documentElement.style.overflow = 'hidden';")

def show_scrollbar(driver):
    driver.execute_script("document.documentElement.style.overflow = '';")

def capture_fullpage_by_scroll(driver, save_path, scroll_pause=0.5):
    """스크롤하며 전체 페이지 캡처"""
    viewport_height = driver.execute_script("return window.innerHeight")
    total_height = driver.execute_script("return document.body.scrollHeight")
    images = []
    driver.execute_script("window.scrollTo(0, 0)")
    time.sleep(scroll_pause)
    y = 0
    while y < total_height:
        driver.execute_script(f"window.scrollTo(0, {y})")
        time.sleep(scroll_pause)
        png = driver.get_screenshot_as_png()
        img = PilImage.open(BytesIO(png))
        images.append(img)
        y += viewport_height
    stitched_height = sum(i.height for i in images)
    result_img = PilImage.new("RGB", (images[0].width, stitched_height), (255, 255, 255)) 
    y_offset = 0
    for img in images:
        result_img.paste(img, (0, y_offset))
        y_offset += img.height
    if total_height < stitched_height:
        result_img = result_img.crop((0, 0, result_img.width, total_height))
    result_img.save(save_path)


# GUI 클래스
class ImageCrawlerGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("통합 이미지 크롤러 (v2: 병렬 처리 방식)")
        self.root.geometry("670x550")
        self.create_widgets()
        self.driver = None

    def create_widgets(self):
        frm = tk.Frame(self.root)
        frm.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)
        row = 0
        tk.Label(frm, text="1. 플랫폼 선택:").grid(row=row, column=0, sticky='w', pady=2)
        self.platform_combo = ttk.Combobox(frm, values=["G마켓", "옥션", "11번가", "쿠팡"], width=48, state="readonly")
        self.platform_combo.grid(row=row, column=1)
        self.platform_combo.current(0)
        row += 1
        tk.Label(frm, text="2. 엑셀 파일 선택:").grid(row=row, column=0, sticky='w', pady=2)
        self.excel_entry = tk.Entry(frm, width=50)
        self.excel_entry.grid(row=row, column=1)
        tk.Button(frm, text="파일 찾기", command=self.select_excel).grid(row=row, column=2)
        row += 1
        tk.Label(frm, text="3. 이미지 저장 폴더:").grid(row=row, column=0, sticky='w', pady=2)
        self.folder_entry = tk.Entry(frm, width=50)
        self.folder_entry.grid(row=row, column=1)
        tk.Button(frm, text="폴더 찾기", command=self.select_folder).grid(row=row, column=2)
        row += 1
        tk.Label(frm, text="4. User-Agent:").grid(row=row, column=0, sticky='w', pady=2)
        self.ua_entry = tk.Entry(frm, width=50)
        self.ua_entry.insert(0, "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36")
        self.ua_entry.grid(row=row, column=1, columnspan=2, sticky='w')
        row += 1
        self.start_btn = tk.Button(frm, text="크롤링 시작", bg="#4CAF50", fg="white", font=("Arial", 10, "bold"), command=self.start_crawl)
        self.start_btn.grid(row=row, column=0, columnspan=3, sticky="ew", pady=10)
        row += 1
        self.progress = ttk.Progressbar(frm, orient="horizontal", length=500, mode="determinate")
        self.progress.grid(row=row, column=0, columnspan=3, pady=5)
        row += 1
        tk.Label(frm, text="로그:").grid(row=row, column=0, sticky='w')
        row += 1
        self.log = scrolledtext.ScrolledText(frm, width=90, height=17, state="normal")
        self.log.grid(row=row, column=0, columnspan=3)

    def select_excel(self):
        path = filedialog.askopenfilename(title="엑셀 파일 선택", filetypes=[("Excel files", "*.xlsx *.xlsm")])
        if path:
            self.excel_entry.delete(0, tk.END)
            self.excel_entry.insert(0, path)

    def select_folder(self):
        folder = filedialog.askdirectory(title="이미지 저장 폴더 선택")
        if folder:
            self.folder_entry.delete(0, tk.END)
            self.folder_entry.insert(0, folder)

    def write_log(self, msg):
        self.log.insert(tk.END, msg)
        self.log.see(tk.END)
        self.root.update_idletasks()

    def start_crawl(self):
        self.start_btn["state"] = "disabled"
        self.log.delete(1.0, tk.END)
        threading.Thread(target=self.crawl_main, daemon=True).start()

    def crawl_main(self):
        try:
            platform = self.platform_combo.get()
            excel_path = self.excel_entry.get()
            save_dir = self.folder_entry.get()
            user_agent = self.ua_entry.get()

            if not os.path.exists(excel_path) or not os.path.exists(save_dir):
                self.write_log("오류: 엑셀 파일과 저장 폴더를 모두 선택하세요.\n")
                self.start_btn["state"] = "normal"
                return

            self.write_log(f"선택된 플랫폼: {platform}\n")
            wb = load_workbook(excel_path)
            ws = wb.active

            url_list = []
            for idx, row in enumerate(ws.iter_rows(min_row=2, min_col=7, max_col=7), start=2):
                cell = row[0]
                url = None
                if cell.hyperlink and cell.hyperlink.target: url = cell.hyperlink.target.strip()
                elif cell.value and isinstance(cell.value, str) and cell.value.strip().startswith("http"): url = cell.value.strip()
                if url: url_list.append((idx, url))
            
            self.write_log(f"총 {len(url_list)}개의 유효한 URL을 읽었습니다.\n")
            self.progress["maximum"] = len(url_list)

            chrome_options = Options()
            chrome_options.page_load_strategy = 'eager'
            chrome_options.add_argument(f'--user-agent={user_agent}')
            # 충돌을 일으키는 불필요한 탐지 우회 옵션 제거
            chrome_options.add_argument("--no-sandbox")
            chrome_options.add_argument("--disable-dev-shm-usage")
            
            self.driver = uc.Chrome(options=chrome_options, use_subprocess=True, version_main=_detect_chrome_major())
            self.driver.set_page_load_timeout(20)
            
            batch_size = 10
            for i in range(0, len(url_list), batch_size):
                batch = url_list[i:i + batch_size]
                tab_mapping = []
                
                try:
                    self.driver.get(batch[0][1])
                except:
                    self.write_log(f"  - (경고) 첫 탭 로드 시간 초과. 다음 작업 계속 진행...\n")
                    pass
                tab_mapping.append((self.driver.current_window_handle, batch[0][0]))

                for row_num, url in batch[1:]:
                    self.driver.switch_to.new_window('tab')
                    try:
                        self.driver.get(url)
                    except:
                        self.write_log(f"  - (경고) [{row_num}]번 탭 로드 시간 초과. 다음 작업 계속 진행...\n")
                        pass
                    new_handle = self.driver.current_window_handle
                    tab_mapping.append((new_handle, row_num))
                    time.sleep(0.5)

                self.write_log(f"\n[{i+1}~{i+len(batch)}]번 URL을 탭으로 열었습니다.\n")

                for handle, row_num in tab_mapping:
                    try:
                        self.driver.switch_to.window(handle)
                        self.write_log(f"[{row_num}]번 URL 처리 시작...\n")

                        if platform == "쿠팡":
                            delay = random.uniform(3.0, 4.0)
                            self.write_log(f"  - (정보) 쿠팡 페이지 로딩을 위해 {delay:.1f}초 대기합니다.\n")
                            time.sleep(delay)
                        else:
                            time.sleep(1)

                        self.write_log(f'  - 상세정보 더보기 버튼 처리 시작...\n')
                        clicked = False
                        if platform == "11번가":
                            if find_and_click_show_more(self.driver): clicked = True
                        elif platform in ["G마켓", "옥션", "쿠팡"]:
                            self.write_log(f"    - {platform} '더보기' 탐색 시작...\n")
                            seemore_xpath = ""
                            if platform in ["G마켓", "옥션"]: seemore_xpath = "//button[.//span[contains(text(), '더보기')]]"
                            elif platform == "쿠팡": seemore_xpath = "/html/body/div[2]/div[1]/main/div[2]/div[4]/div[4]/div[2]"
                            try:
                                self.write_log("      - 1단계: 메인 페이지에서 탐색...\n")
                                wait = WebDriverWait(self.driver, 2)
                                button = wait.until(EC.presence_of_element_located((By.XPATH, seemore_xpath)))
                                self.driver.execute_script("arguments[0].scrollIntoView({block: 'center', inline: 'nearest'});", button)
                                time.sleep(0.3)
                                self.driver.execute_script("arguments[0].click();", button)
                                clicked = True
                                self.write_log("      - 성공! 메인 페이지에서 버튼 클릭 완료.\n")
                            except Exception:
                                self.write_log("      - 1단계 실패. 2단계: iframe 내부 탐색 시작...\n")
                                try:
                                    iframes = self.driver.find_elements(By.TAG_NAME, 'iframe')
                                    if not iframes: self.write_log("      - iframe이 없어 탐색을 종료합니다.\n")
                                    else:
                                        self.write_log(f"      - {len(iframes)}개의 iframe 발견. 순차적으로 탐색합니다.\n")
                                        for i_frame, iframe in enumerate(iframes):
                                            try:
                                                self.write_log(f"        - {i_frame + 1}/{len(iframes)}번 iframe 탐색 중...\n")
                                                self.driver.switch_to.frame(iframe)
                                                wait_iframe = WebDriverWait(self.driver, 1)
                                                button = wait_iframe.until(EC.element_to_be_clickable((By.XPATH, seemore_xpath)))
                                                self.driver.execute_script("arguments[0].click();", button)
                                                clicked = True
                                                self.write_log(f"        - 성공! iframe에서 버튼 클릭 완료.\n")
                                                break 
                                            except Exception: continue
                                            finally: self.driver.switch_to.default_content()
                                        if clicked: break
                                except Exception: self.write_log("      - 2단계 iframe 탐색 중 오류 발생.\n")
                        if clicked: self.write_log(f"  - 상세정보 더보기 처리 완료!\n")
                        else: self.write_log(f"  - (정보) 더보기 버튼이 없거나 클릭하지 않고 다음으로 진행합니다.\n")
                        time.sleep(1)

                        self.write_log(f'  - 이미지 로딩을 위해 3초간 스크롤합니다.\n')
                        end_time = time.time() + 3
                        while time.time() < end_time:
                            self.driver.find_element(By.TAG_NAME, 'body').send_keys(Keys.END)
                            time.sleep(0.2)
                        
                        self.write_log(f'  - 전체 페이지 캡처 시작...\n')
                        selectors_to_hide = []
                        if platform == "G마켓": selectors_to_hide = ['#desktop_layout-header']
                        elif platform == "옥션": selectors_to_hide = ['#ulVipTab', '#header']
                        elif platform == "11번가": selectors_to_hide = ['#header']
                        elif platform == "쿠팡": selectors_to_hide = ['#header', '.wing-banner']
                        start_hide_fixed_ui_elements_interval(self.driver, selectors_to_hide, interval=200)

                        xpaths_to_hide = []
                        if platform == "G마켓": xpaths_to_hide = ['//*[@id="container"]/div[7]/div[1]','//*[@id="container"]/div[7]/div[4]/div','//*[@id="desktop_layout-footer"]/div[2]/a/span']
                        elif platform == "옥션": xpaths_to_hide = ['//*[@id="content"]/div[4]/div[3]/div', '//*[@id="btnSkipTop"]']
                        elif platform == "11번가": xpaths_to_hide = ['//*[@id="gnb"]/div/div[2]','//*[@id="layBodyWrap"]/div/div[1]/div[2]/div/div[2]/div[2]','//*[@id="footer"]/div[1]/div','//*[@id="tabDetail"]/ul','//*[@id="layBodyWrap"]/div/div[1]/div[2]/div/div[2]']
                        elif platform == "쿠팡": xpaths_to_hide = ['/html/body/div[2]/div[1]/main/div[2]/div[3]','//*[@id="wa-goto-top"]']
                        hide_elements_by_xpath(self.driver, xpaths_to_hide)

                        hide_scrollbar(self.driver)
                        time.sleep(0.6)

                        today = datetime.datetime.now().strftime("%Y%m%d")
                        product_code = ws[f'H{row_num}'].value or f"NO_CODE_{row_num}"
                        save_path = os.path.join(save_dir, f"{product_code}_{today}.png")
                        capture_fullpage_by_scroll(self.driver, save_path, scroll_pause=0.6)
                        stop_hide_fixed_ui_elements_interval(self.driver)
                        show_scrollbar(self.driver)
                        self.write_log(f'  - 캡처 저장 완료: {os.path.basename(save_path)}\n')
                    except Exception as e:
                        self.write_log(f"[{row_num}] 처리 중 오류 발생: {e}\n")
                        traceback.print_exc()
                    
                    self.progress["value"] += 1
                    self.root.update_idletasks()

                main_handle = tab_mapping[0][0]
                # --- [unpack 오류 수정] ---
                for handle, _ in tab_mapping[1:]:
                    try:
                        self.driver.switch_to.window(handle)
                        self.driver.close()
                    except: pass
                self.driver.switch_to.window(main_handle)
                self.write_log("이번 배치 탭 정리 완료.\n")

            messagebox.showinfo("작업 완료", "모든 URL 크롤링 및 캡처가 완료되었습니다!")
        except Exception as e:
            err_msg = str(e)
            if "session not created" in err_msg.lower() or "DevToolsActivePort" in err_msg:
                version_msg = "오류: Chrome 브라우저와 ChromeDriver 버전이 맞지 않습니다."
                self.write_log(version_msg + "\n")
                messagebox.showerror("ChromeDriver 버전 오류", version_msg)
            else:
                self.write_log(f"전체 처리 중 예외 발생: {e}\n")
                traceback.print_exc()
        finally:
            if self.driver: self.driver.quit()
            self.write_log("\n브라우저 종료 및 모든 작업 완료!\n")
            self.start_btn["state"] = "normal"

if __name__ == "__main__":
    root = tk.Tk()
    app = ImageCrawlerGUI(root)
    root.mainloop()