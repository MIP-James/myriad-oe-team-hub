import os
import threading
import queue
import traceback
import time
from dataclasses import dataclass
from typing import Any, List, Optional, Tuple

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait, Select as SeleniumSelect
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoAlertPresentException


def _detect_chrome_major():
    """Registry 에서 설치된 Chrome major 버전 읽기. UC version_main=None 이 CfT 최신 드라이버를 받아오는 동작 회피용. 실패 시 None → UC 기본 동작."""
    try:
        import winreg
        for hive, path in (
            (winreg.HKEY_CURRENT_USER, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Wow6432Node\Google\Chrome\BLBeacon"),
        ):
            try:
                with winreg.OpenKey(hive, path) as key:
                    version, _ = winreg.QueryValueEx(key, "version")
                    return int(str(version).split(".")[0])
            except OSError:
                continue
    except Exception:
        pass
    return None


APP_TITLE = "Auction VeRO 신고 매크로"
VERO_URL = "https://memberssl.auction.co.kr/vero/whatisVero.aspx"
BATCH_TAB_COUNT = 10

YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF59D")
GREEN_FILL = PatternFill(fill_type="solid", fgColor="C8E6C9")
RED_FILL = PatternFill(fill_type="solid", fgColor="FFCDD2")

# G마켓 템플릿 침해유형 -> 옥션 VeRO 신고 사유 매핑
REPORT_TYPE_MAP = {
    "위조품": "1",
    "위조상품": "1",
    "상표권 침해": "6",
    "상표권침해": "6",
    "저작권 침해": "5",
    "저작권침해": "5",
    "디자인권 침해": "4",
    "디자인권침해": "4",
}


@dataclass
class ReportItem:
    row_no: int
    product_no: str
    report_type_text: str
    product_url: str
    status: str = "대기"
    detail: str = ""
    result_fill: Optional[str] = None


class Selectors:
    LOGIN_ID = [
        (By.CSS_SELECTOR, "input#VeroLogin_htxtVeroID"),
        (By.CSS_SELECTOR, 'input[name="VeroLogin$htxtVeroID"]'),
    ]
    LOGIN_PW = [
        (By.CSS_SELECTOR, "input#VeroLogin_htxtPassWord"),
        (By.CSS_SELECTOR, 'input[name="VeroLogin$htxtPassWord"]'),
        (By.CSS_SELECTOR, 'input[type="password"]'),
    ]
    LOGIN_SUBMIT = [
        (By.CSS_SELECTOR, "input#VeroLogin_ibtnVeroLogIn"),
        (By.CSS_SELECTOR, 'input[name="VeroLogin$ibtnVeroLogIn"]'),
        (By.XPATH, '//div[contains(@class,"login")]//input[@type="image"]'),
    ]
    LOGIN_SUCCESS = [
        (By.CSS_SELECTOR, "input#VeroLogin_ibtnVeroLogOut"),
        (By.CSS_SELECTOR, 'input[name="VeroLogin$ibtnVeroLogOut"]'),
        (By.XPATH, '//p[contains(@class,"uxb-welcome")]'),
    ]

    ITEM_STOPPED = [
        (By.XPATH, '//img[@alt="죄송합니다. 현재 진행중인 상품이 없습니다."]'),
        (By.XPATH, '//img[contains(@src,"txt_infos04.gif")]'),
    ]
    EXCHANGE_TAB = [
        (By.XPATH, '//a[@data-maintarget="#vip_tab_exchange"]'),
        (By.XPATH, '//a[contains(@href,"#vip_tab_exchange")]'),
        (By.XPATH, '//a[normalize-space()="교환/반품"]'),
    ]
    VERO_SECTION = [
        (By.CSS_SELECTOR, "#ucReport_hdivVeroRegist"),
        (By.XPATH, '//div[@id="ucReport_hdivVeroRegist"]'),
    ]
    VERO_REASON_SELECT = (By.CSS_SELECTOR, "select#repveroreason_select")
    VERO_SUBMIT = [
        (By.XPATH, '//img[@title="VeRO 신고하기"]'),
        (By.XPATH, '//img[@alt="VeRO 신고하기"]'),
        (By.XPATH, '//img[contains(@onclick,"VeroRegist.Check")]'),
    ]


def normalize_text(text: str) -> str:
    return (text or "").strip().replace(" ", "").replace("\u00a0", "")


def get_report_type_value(label: str) -> Optional[str]:
    n = normalize_text(label)
    for k, v in REPORT_TYPE_MAP.items():
        if normalize_text(k) == n:
            return v
    return None


def safe_str(value) -> str:
    return "" if value is None else str(value).strip()


def reset_frame(driver: uc.Chrome):
    try:
        driver.switch_to.default_content()
    except Exception:
        pass


def get_frames(driver: uc.Chrome):
    frames = []
    for tag in ("frame", "iframe"):
        try:
            frames.extend(driver.find_elements(By.TAG_NAME, tag))
        except Exception:
            pass
    return frames


def get_iframe_srcs(driver: uc.Chrome) -> List[Tuple[int, str]]:
    reset_frame(driver)
    srcs = []
    for idx, frame in enumerate(get_frames(driver)):
        try:
            srcs.append((idx, frame.get_attribute("src") or ""))
        except Exception:
            srcs.append((idx, ""))
    return srcs


def switch_to_frame_idx(driver: uc.Chrome, idx: int):
    reset_frame(driver)
    if idx >= 0:
        frames = get_frames(driver)
        if idx < len(frames):
            driver.switch_to.frame(frames[idx])


def find_element_in_frames(
    driver: uc.Chrome,
    by: str,
    selector: str,
    visible: bool = True,
    timeout: float = 1.2,
) -> Tuple[Optional[int], Any]:
    cond = EC.visibility_of_element_located if visible else EC.presence_of_element_located

    def _try(idx: int):
        try:
            el = WebDriverWait(driver, timeout).until(cond((by, selector)))
            return idx, el
        except Exception:
            return None, None

    reset_frame(driver)
    fi, el = _try(-1)
    if el is not None:
        return fi, el

    for i, frame in enumerate(get_frames(driver)):
        reset_frame(driver)
        try:
            driver.switch_to.frame(frame)
        except Exception:
            continue
        fi, el = _try(i)
        if el is not None:
            return fi, el

    reset_frame(driver)
    return None, None


def find_first_in_frames(
    driver: uc.Chrome,
    candidates: List[Tuple[str, str]],
    visible: bool = True,
    timeout: float = 1.2,
) -> Tuple[Optional[int], Any, Optional[Tuple[str, str]]]:
    for by, selector in candidates:
        fi, el = find_element_in_frames(driver, by, selector, visible=visible, timeout=timeout)
        if el is not None:
            return fi, el, (by, selector)
    return None, None, None


def accept_alert(driver: uc.Chrome, timeout: float = 2.0) -> Optional[str]:
    try:
        alert = WebDriverWait(driver, timeout).until(EC.alert_is_present())
        msg = alert.text
        alert.accept()
        return msg
    except Exception:
        return None


def safe_fill(element, text: str):
    try:
        element.click()
    except Exception:
        pass
    try:
        element.clear()
    except Exception:
        pass
    element.send_keys(text)


class AuctionVeroReporter:
    def __init__(self, user_id: str, password: str, log_callback=None):
        self.user_id = user_id
        self.password = password
        self.log_callback = log_callback or (lambda m: None)
        self.driver: Optional[uc.Chrome] = None
        self.login_handle: Optional[str] = None

    def log(self, msg: str):
        self.log_callback(msg)

    def start(self):
        opts = uc.ChromeOptions()
        opts.add_argument("--window-size=1440,1200")
        opts.add_argument("--no-sandbox")
        opts.add_argument("--disable-dev-shm-usage")
        # 옥션 상품 페이지는 광고/트래킹 스크립트 때문에 load 이벤트가 40초를 넘기는 경우가 있음
        # → DOM 준비(DOMContentLoaded)까지만 기다리고 진행 (Image_Crawler 와 동일)
        opts.page_load_strategy = "eager"
        self.driver = uc.Chrome(
            headless=False,
            options=opts,
            use_subprocess=True,  # frozen launcher 환경에서 통합 런처 재실행 방지
            version_main=_detect_chrome_major(),  # CfT 최신 드라이버 자동 다운로드 회피
        )
        self.driver.implicitly_wait(0)
        self.driver.set_page_load_timeout(40)

    def close(self):
        try:
            if self.driver:
                self.driver.quit()
        except Exception:
            pass

    def open_login_page(self):
        self.log("[접속] Auction VeRO 페이지 접속")
        self.navigate(VERO_URL)
        self.login_handle = self.driver.current_window_handle
        time.sleep(2)

    def login(self):
        self.open_login_page()

        for i, src in get_iframe_srcs(self.driver):
            self.log(f"[프레임] {i}: {src}")

        fi_id, el_id, sel_id = find_first_in_frames(self.driver, Selectors.LOGIN_ID, timeout=2.0)
        fi_pw, el_pw, sel_pw = find_first_in_frames(self.driver, Selectors.LOGIN_PW, timeout=2.0)

        if el_id is None:
            raise RuntimeError("옥션 VeRO ID 입력칸을 찾지 못했습니다.")
        if el_pw is None:
            raise RuntimeError("옥션 VeRO PW 입력칸을 찾지 못했습니다.")

        self.log(f"[로그인] ID selector={sel_id} frame_idx={fi_id}")
        self.log(f"[로그인] PW selector={sel_pw} frame_idx={fi_pw}")

        switch_to_frame_idx(self.driver, fi_id if fi_id is not None else -1)
        _, el_id2, _ = find_first_in_frames(self.driver, Selectors.LOGIN_ID, timeout=1.0)
        if el_id2 is None:
            raise RuntimeError("ID 입력칸 재탐색 실패")
        safe_fill(el_id2, self.user_id)

        switch_to_frame_idx(self.driver, fi_pw if fi_pw is not None else -1)
        fi_pw2, el_pw2, _ = find_first_in_frames(self.driver, Selectors.LOGIN_PW, timeout=1.0)
        if el_pw2 is None:
            raise RuntimeError("PW 입력칸 재탐색 실패")
        safe_fill(el_pw2, self.password)

        fi_btn, el_btn, sel_btn = find_first_in_frames(self.driver, Selectors.LOGIN_SUBMIT, timeout=2.0)
        self.log(f"[로그인] BTN selector={sel_btn} frame_idx={fi_btn}")

        submit_errors = []
        submitted = False

        if el_btn is not None:
            try:
                switch_to_frame_idx(self.driver, fi_btn if fi_btn is not None else -1)
                _, el_btn2, _ = find_first_in_frames(self.driver, Selectors.LOGIN_SUBMIT, timeout=1.0)
                if el_btn2 is not None:
                    try:
                        el_btn2.click()
                    except Exception:
                        self.driver.execute_script("arguments[0].click();", el_btn2)
                    submitted = True
                    self.log("[로그인] 버튼 클릭 완료")
            except Exception as e:
                submit_errors.append(f"버튼 클릭 실패: {e}")

        if not submitted:
            try:
                switch_to_frame_idx(self.driver, fi_btn if fi_btn is not None else (fi_pw2 if fi_pw2 is not None else -1))
                self.driver.execute_script("if (typeof fnLogin2 === 'function') { return fnLogin2(); }")
                submitted = True
                self.log("[로그인] fnLogin2() JS 실행 완료")
            except Exception as e:
                submit_errors.append(f"fnLogin2() 실패: {e}")

        if not submitted:
            try:
                switch_to_frame_idx(self.driver, fi_pw2 if fi_pw2 is not None else -1)
                _, el_pw3, _ = find_first_in_frames(self.driver, Selectors.LOGIN_PW, timeout=1.0)
                if el_pw3 is not None:
                    el_pw3.send_keys(Keys.RETURN)
                    submitted = True
                    self.log("[로그인] Enter 제출 완료")
            except Exception as e:
                submit_errors.append(f"Enter 실패: {e}")

        alert_msg = accept_alert(self.driver, timeout=2.0)
        if alert_msg:
            self.log(f"[로그인] alert: {alert_msg}")

        time.sleep(3)

        success_used = None
        fi_s, el_s, sel_s = find_first_in_frames(self.driver, Selectors.LOGIN_SUCCESS, visible=True, timeout=3.0)
        if el_s is not None:
            success_used = f"{sel_s} (로그아웃 버튼/환영문구 기준, frame_idx={fi_s})"

        if not success_used:
            fi_s, el_s, sel_s = find_first_in_frames(self.driver, Selectors.LOGIN_SUCCESS, visible=False, timeout=1.5)
            if el_s is not None:
                success_used = f"{sel_s} (attached, frame_idx={fi_s})"

        current_url = self.driver.current_url
        self.log(f"[로그인후 URL] {current_url}")

        if not success_used:
            joined = " | ".join(submit_errors) if submit_errors else "제출 시도 기록 없음"
            raise RuntimeError(
                "옥션 VeRO 로그인 성공을 확인하지 못했습니다. "
                f"제출 로그: {joined}"
            )

        self.log(f"[로그인] 성공 확인: {success_used}")
        self.login_handle = self.driver.current_window_handle

    def navigate(self, url: str):
        """페이지 이동. 로딩 타임아웃은 실패로 보지 않음 — 페이지는 백그라운드에서 계속 로딩되고,
        실제 신고 처리 시점에 다시 확인함 (report_current_item 에서 1회 새로고침 재시도)."""
        try:
            self.driver.get(url)
        except TimeoutException:
            self.log(f"[탭] 페이지 로딩 지연 → 계속 진행 (처리 시 재확인): {url}")

    def open_batch_tabs(self, batch_items: List[ReportItem]) -> List[Tuple[ReportItem, str]]:
        pairs: List[Tuple[ReportItem, str]] = []

        for item in batch_items:
            self.driver.switch_to.window(self.login_handle)
            handle = None
            open_errors = []

            # 1) Selenium native 새 탭 생성
            #    (탭 생성에 성공했으면 로딩이 느려도 그 탭을 그대로 사용 — 예전엔 로딩 타임아웃을
            #     탭 생성 실패로 오인해 로그인 탭까지 덮어쓰다 전체가 중단됐음)
            try:
                self.driver.switch_to.new_window("tab")
                handle = self.driver.current_window_handle
                self.log(f"[탭] selenium new_window 성공 / handle={handle}")
            except Exception as e:
                open_errors.append(f"new_window 실패: {e}")
                handle = None
            if handle is not None:
                self.navigate(item.product_url)

            # 2) 실패 시 빈 탭 JS 생성 후 이동
            if handle is None:
                try:
                    self.driver.switch_to.window(self.login_handle)
                    before_handles = list(self.driver.window_handles)
                    before_set = set(before_handles)
                    self.driver.execute_script("window.open('about:blank', '_blank');")
                    time.sleep(1.0)
                    after_handles = list(self.driver.window_handles)
                    new_handles = [h for h in after_handles if h not in before_set]
                    if new_handles:
                        handle = new_handles[-1]
                        self.driver.switch_to.window(handle)
                        self.log(f"[탭] JS blank tab 성공 / handle={handle}")
                        self.navigate(item.product_url)
                except Exception as e:
                    open_errors.append(f"blank_open 실패: {e}")
                    handle = None

            # 3) 마지막 수단: 현재 탭에서 열고, 로그인 탭은 마지막에 재포커스
            if handle is None:
                try:
                    self.driver.switch_to.window(self.login_handle)
                    self.driver.execute_script("window.location.href = arguments[0];", item.product_url)
                    time.sleep(1.0)
                    handle = self.driver.current_window_handle
                    self.log(f"[탭] 현재 탭 이동 fallback 사용 / handle={handle}")
                except Exception as e:
                    open_errors.append(f"same_tab_fallback 실패: {e}")
                    handle = None

            if handle is None:
                joined = " | ".join(open_errors) if open_errors else "원인 미상"
                raise RuntimeError(f"새 탭 생성 실패: {item.product_url} / {joined}")

            pairs.append((item, handle))
            self.log(f"[탭] 열기: 행 {item.row_no} / 상품번호 {item.product_no} / handle={handle}")

        return pairs

    def is_stopped_item_page(self) -> bool:
        for by, selector in Selectors.ITEM_STOPPED:
            try:
                elems = self.driver.find_elements(by, selector)
                if elems:
                    return True
            except Exception:
                continue
        return False

    def goto_exchange_tab(self):
        fi, el, used = find_first_in_frames(self.driver, Selectors.EXCHANGE_TAB, visible=True, timeout=4.0)
        if el is None:
            raise RuntimeError("교환/반품 버튼을 찾지 못했습니다.")
        switch_to_frame_idx(self.driver, fi if fi is not None else -1)
        _, el2, _ = find_first_in_frames(self.driver, Selectors.EXCHANGE_TAB, visible=True, timeout=1.0)
        if el2 is None:
            raise RuntimeError("교환/반품 버튼 재탐색 실패")
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el2)
        except Exception:
            pass
        time.sleep(0.5)
        try:
            el2.click()
        except Exception:
            self.driver.execute_script("arguments[0].click();", el2)
        self.log(f"[상품] 교환/반품 클릭 완료 / selector={used}")
        time.sleep(1.2)

    def goto_vero_section(self):
        try:
            self.driver.execute_script("""
                var el = document.getElementById('ucReport_hdivVeroRegist');
                if (el) { el.scrollIntoView({block:'center'}); return true; }
                return false;
            """)
        except Exception:
            pass

        fi, el, _ = find_first_in_frames(self.driver, Selectors.VERO_SECTION, visible=True, timeout=4.0)
        if el is None:
            fi, el, _ = find_first_in_frames(self.driver, Selectors.VERO_SECTION, visible=False, timeout=2.0)
        if el is None:
            raise RuntimeError("[VeRO 신고] 영역을 찾지 못했습니다.")
        switch_to_frame_idx(self.driver, fi if fi is not None else -1)
        time.sleep(0.4)

    def select_report_type(self, report_type_text: str):
        value = get_report_type_value(report_type_text)
        if not value:
            supported = ", ".join(sorted(REPORT_TYPE_MAP.keys()))
            raise RuntimeError(
                f"지원하지 않는 침해유형: '{report_type_text}' / 지원값: {supported}"
            )

        fi, el = find_element_in_frames(self.driver, *Selectors.VERO_REASON_SELECT, visible=True, timeout=2.0)
        if el is None:
            fi, el = find_element_in_frames(self.driver, *Selectors.VERO_REASON_SELECT, visible=False, timeout=1.0)
        if el is None:
            raise RuntimeError("VeRO 신고 사유 드롭다운을 찾지 못했습니다.")

        switch_to_frame_idx(self.driver, fi if fi is not None else -1)
        _, el2 = find_element_in_frames(self.driver, *Selectors.VERO_REASON_SELECT, visible=False, timeout=1.0)
        if el2 is None:
            raise RuntimeError("VeRO 신고 사유 드롭다운 재탐색 실패")

        SeleniumSelect(el2).select_by_value(value)
        self.log(f"[신고] 침해유형 선택: {report_type_text} ({value})")

    def submit_report(self):
        fi, el, used = find_first_in_frames(self.driver, Selectors.VERO_SUBMIT, visible=True, timeout=2.0)
        if el is None:
            fi, el, used = find_first_in_frames(self.driver, Selectors.VERO_SUBMIT, visible=False, timeout=1.0)
        if el is None:
            raise RuntimeError("VeRO 신고하기 버튼을 찾지 못했습니다.")

        switch_to_frame_idx(self.driver, fi if fi is not None else -1)
        _, el2, _ = find_first_in_frames(self.driver, Selectors.VERO_SUBMIT, visible=False, timeout=1.0)
        if el2 is None:
            raise RuntimeError("VeRO 신고하기 버튼 재탐색 실패")

        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el2)
        except Exception:
            pass
        time.sleep(0.3)

        try:
            el2.click()
        except Exception:
            self.driver.execute_script("arguments[0].click();", el2)
        self.log(f"[신고] 버튼 클릭 완료 / selector={used}")

        alert_msg = accept_alert(self.driver, timeout=5.0)
        if not alert_msg:
            raise RuntimeError("신고 완료 팝업을 확인하지 못했습니다.")
        self.log(f"[신고] alert: {alert_msg}")

        if "침해물품으로 신고 되었습니다" not in alert_msg:
            raise RuntimeError(f"예상과 다른 팝업 메시지: {alert_msg}")

    def wait_page_ready(self):
        try:
            WebDriverWait(self.driver, 10).until(
                lambda d: d.execute_script("return document.readyState") == "complete"
            )
        except Exception:
            pass
        time.sleep(1.0)

    def report_current_item(self, item: ReportItem) -> Tuple[str, str]:
        self.log(f"[상품] URL 접속 확인: {item.product_url}")
        self.wait_page_ready()

        if self.is_stopped_item_page():
            return "진행 불가", "판매 중단된 상품 또는 현재 진행중인 상품 없음"

        try:
            self.goto_exchange_tab()
        except Exception as e:
            # 탭을 열 때 로딩이 지연돼 페이지가 덜 그려졌을 수 있음 → 1회 새로고침 후 재시도
            self.log(f"[상품] {e} → 페이지 다시 열어 재시도")
            self.navigate(item.product_url)
            self.wait_page_ready()
            if self.is_stopped_item_page():
                return "진행 불가", "판매 중단된 상품 또는 현재 진행중인 상품 없음"
            self.goto_exchange_tab()
        self.goto_vero_section()
        self.select_report_type(item.report_type_text)
        self.submit_report()
        return "신고 완료", "정상 처리됨"

    def process_batch(self, batch_items: List[ReportItem]) -> List[Tuple[int, str, str]]:
        results: List[Tuple[int, str, str]] = []
        tab_pairs = self.open_batch_tabs(batch_items)

        for item, handle in tab_pairs:
            try:
                self.driver.switch_to.window(handle)
                status, detail = self.report_current_item(item)
                results.append((item.row_no, status, detail))
            except TimeoutException as e:
                results.append((item.row_no, "실패", f"타임아웃: {e}"))
            except Exception as e:
                results.append((item.row_no, "실패", str(e)))
            finally:
                try:
                    current_handles = self.driver.window_handles
                    if handle != self.login_handle and handle in current_handles:
                        self.driver.close()
                        self.driver.switch_to.window(self.login_handle)
                    else:
                        # same-tab fallback이 사용된 경우 로그인 페이지를 다시 열어 세션 유지 상태 복귀
                        self.navigate(VERO_URL)
                        time.sleep(1.5)
                        self.login_handle = self.driver.current_window_handle
                        fi_s, el_s, _ = find_first_in_frames(self.driver, Selectors.LOGIN_SUCCESS, visible=True, timeout=2.0)
                        if el_s is None:
                            fi_s, el_s, _ = find_first_in_frames(self.driver, Selectors.LOGIN_SUCCESS, visible=False, timeout=1.0)
                        if el_s is None:
                            raise RuntimeError("상품 탭 처리 후 VeRO 로그인 세션 복귀 확인 실패")
                except Exception as close_err:
                    self.log(f"[탭] 정리 중 경고: {close_err}")
                    try:
                        if self.login_handle in self.driver.window_handles:
                            self.driver.switch_to.window(self.login_handle)
                    except Exception:
                        pass
                time.sleep(0.2)

        return results


class LoginDialog(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("Auction VeRO 로그인 정보 입력")
        self.resizable(False, False)
        self.result = None
        self.transient(parent)
        self.grab_set()

        frame = ttk.Frame(self, padding=12)
        frame.pack(fill="both", expand=True)

        ttk.Label(frame, text="ID").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Label(frame, text="PW").grid(row=1, column=0, sticky="w", pady=4)

        self.id_var = tk.StringVar()
        self.pw_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.id_var, width=30).grid(row=0, column=1, pady=4)
        ttk.Entry(frame, textvariable=self.pw_var, width=30, show="*").grid(row=1, column=1, pady=4)

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=2, column=0, columnspan=2, pady=(12, 0))
        ttk.Button(btn_frame, text="확인", command=self.on_ok).pack(side="left", padx=4)
        ttk.Button(btn_frame, text="취소", command=self.on_cancel).pack(side="left", padx=4)

        self.bind("<Return>", lambda e: self.on_ok())
        self.bind("<Escape>", lambda e: self.on_cancel())
        self.update_idletasks()
        w = self.winfo_width()
        h = self.winfo_height()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = int((sw / 2) - (w / 2))
        y = int((sh / 2) - (h / 2))
        self.geometry(f"+{x}+{y}")

    def on_ok(self):
        uid = self.id_var.get().strip()
        pw = self.pw_var.get().strip()
        if not uid or not pw:
            messagebox.showwarning("입력 필요", "ID와 PW를 모두 입력해주세요.", parent=self)
            return
        self.result = (uid, pw)
        self.destroy()

    def on_cancel(self):
        self.result = None
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("980x760")
        self.center_window(980, 760)

        self.items: List[ReportItem] = []
        self.excel_path: Optional[str] = None
        self.log_queue = queue.Queue()
        self.worker_thread: Optional[threading.Thread] = None
        self.progress_var = tk.DoubleVar(value=0)
        self.progress_text_var = tk.StringVar(value="진행률: 0 / 0")

        self._build_ui()
        self.after(150, self._poll_log_queue)

    def center_window(self, width, height):
        self.update_idletasks()
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        x = int((screen_width / 2) - (width / 2))
        y = int((screen_height / 2) - (height / 2))
        self.geometry(f"{width}x{height}+{x}+{y}")

    def _build_ui(self):
        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")

        ttk.Button(top, text="엑셀 업로드", command=self.load_excel).pack(side="left", padx=4)
        ttk.Button(top, text="신고 진행", command=self.start_report).pack(side="left", padx=4)
        ttk.Button(top, text="결과 엑셀 저장", command=self.export_result).pack(side="left", padx=4)
        self.file_label = ttk.Label(top, text="업로드된 파일 없음")
        self.file_label.pack(side="left", padx=10)

        progress_frame = ttk.Frame(self, padding=(10, 0, 10, 8))
        progress_frame.pack(fill="x")
        self.progress_bar = ttk.Progressbar(
            progress_frame,
            variable=self.progress_var,
            maximum=100,
            mode="determinate",
        )
        self.progress_bar.pack(fill="x", side="left", expand=True)
        ttk.Label(progress_frame, textvariable=self.progress_text_var, width=18).pack(side="left", padx=(10, 0))

        mid = ttk.Frame(self, padding=(10, 0, 10, 0))
        mid.pack(fill="both", expand=True)

        cols = ("row_no", "product_no", "report_type", "product_url", "status", "detail")
        self.tree = ttk.Treeview(mid, columns=cols, show="headings", height=14)
        hdrs = dict(
            row_no="엑셀행",
            product_no="상품번호",
            report_type="침해유형",
            product_url="상품 URL",
            status="상태",
            detail="상세결과",
        )
        widths = dict(row_no=70, product_no=120, report_type=140, product_url=310, status=100, detail=280)
        for c in cols:
            self.tree.heading(c, text=hdrs[c])
            self.tree.column(c, width=widths[c], anchor="w")

        vsb = ttk.Scrollbar(mid, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(mid, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        mid.grid_rowconfigure(0, weight=1)
        mid.grid_columnconfigure(0, weight=1)

        bottom = ttk.LabelFrame(self, text="로그", padding=10)
        bottom.pack(fill="both", expand=False, padx=10, pady=10)

        self.log_text = tk.Text(bottom, height=16, wrap="word")
        self.log_text.pack(fill="both", expand=True)

    def log(self, msg: str):
        self.log_queue.put(msg)

    def _poll_log_queue(self):
        while not self.log_queue.empty():
            msg = self.log_queue.get_nowait()
            self.log_text.insert("end", msg + "\n")
            self.log_text.see("end")
        self.after(150, self._poll_log_queue)

    def reset_progress(self):
        total = len(self.items)
        self.progress_var.set(0)
        self.progress_text_var.set(f"진행률: 0 / {total}")

    def update_progress(self, completed: int, total: int):
        percent = (completed / total * 100) if total else 0
        self.progress_var.set(percent)
        self.progress_text_var.set(f"진행률: {completed} / {total}")

    def refresh_tree(self):
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        for item in self.items:
            self.tree.insert(
                "",
                "end",
                values=(item.row_no, item.product_no, item.report_type_text, item.product_url, item.status, item.detail),
            )

    def load_excel(self):
        path = filedialog.askopenfilename(
            title="엑셀 템플릿 선택",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if not path:
            return

        try:
            wb = load_workbook(path)
            ws = wb.active

            items: List[ReportItem] = []
            row = 2
            while True:
                product_no = safe_str(ws[f"A{row}"].value)
                report_type_text = safe_str(ws[f"B{row}"].value)
                product_url = safe_str(ws[f"D{row}"].value)

                if not product_no and not report_type_text and not product_url:
                    break

                if product_no:
                    items.append(
                        ReportItem(
                            row_no=row,
                            product_no=product_no,
                            report_type_text=report_type_text,
                            product_url=product_url,
                        )
                    )
                row += 1

            self.items = items
            self.excel_path = path
            self.file_label.config(text=os.path.basename(path))
            self.refresh_tree()
            self.reset_progress()
            self.log(f"[엑셀] {len(items)}건 로드 완료")
            messagebox.showinfo("완료", f"{len(items)}건을 불러왔습니다.")
        except Exception as e:
            messagebox.showerror("오류", f"엑셀 로드 실패\n{e}")

    def update_item_result(self, index: int, status: str, detail: str):
        self.items[index].status = status
        self.items[index].detail = detail

        if status == "신고 완료":
            self.items[index].result_fill = "green"
        elif status == "진행 불가":
            self.items[index].result_fill = "yellow"
        else:
            self.items[index].result_fill = "red"

        self.after(0, self.refresh_tree)

    def start_report(self):
        if not self.items:
            messagebox.showwarning("대상 없음", "먼저 엑셀 파일을 업로드해주세요.")
            return

        if self.worker_thread and self.worker_thread.is_alive():
            messagebox.showwarning("진행 중", "이미 작업이 진행 중입니다.")
            return

        dlg = LoginDialog(self)
        self.wait_window(dlg)
        if not dlg.result:
            return

        user_id, pw = dlg.result

        for item in self.items:
            item.status = "대기"
            item.detail = ""
            item.result_fill = None
        self.refresh_tree()
        self.reset_progress()

        self.worker_thread = threading.Thread(
            target=self._run_report_worker,
            args=(user_id, pw),
            daemon=True,
        )
        self.worker_thread.start()

    def _run_report_worker(self, user_id: str, pw: str):
        reporter = AuctionVeroReporter(user_id, pw, log_callback=self.log)
        success = skipped = failed = 0

        try:
            reporter.start()
            reporter.login()

            total = len(self.items)
            completed = 0

            for batch_start in range(0, total, BATCH_TAB_COUNT):
                batch_items = self.items[batch_start:batch_start + BATCH_TAB_COUNT]
                self.log(f"[배치] 시작: {batch_start + 1} ~ {batch_start + len(batch_items)} / 총 {total}건")

                results = reporter.process_batch(batch_items)
                result_map = {row_no: (status, detail) for row_no, status, detail in results}

                for idx in range(batch_start, batch_start + len(batch_items)):
                    item = self.items[idx]
                    status, detail = result_map.get(item.row_no, ("실패", "결과 누락"))
                    self.update_item_result(idx, status, detail)

                    if status == "신고 완료":
                        success += 1
                    elif status == "진행 불가":
                        skipped += 1
                    else:
                        failed += 1

                    completed += 1
                    self.after(0, lambda c=completed, t=total: self.update_progress(c, t))
                    self.log(f"[처리] 결과: {item.product_no} -> {status} / {detail}")

            summary = (
                f"전체 {len(self.items)}건\n"
                f"신고 완료: {success}건\n"
                f"진행 불가: {skipped}건\n"
                f"실패: {failed}건"
            )
            self.log("[완료] " + summary.replace("\n", " / "))
            self.after(0, lambda: messagebox.showinfo("신고 작업 완료", summary))

        except Exception as e:
            err_msg = str(e)
            self.log("[오류] " + err_msg)
            self.log(traceback.format_exc())
            self.after(0, lambda m=err_msg: messagebox.showerror("실행 오류", m))
        finally:
            reporter.close()

    def export_result(self):
        if not self.items:
            messagebox.showwarning("대상 없음", "먼저 엑셀을 업로드해주세요.")
            return

        if not self.excel_path:
            messagebox.showwarning("원본 없음", "원본 엑셀 경로를 찾을 수 없습니다.")
            return

        save_path = filedialog.asksaveasfilename(
            title="결과 엑셀 저장",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="auction_vero_report_result.xlsx",
        )
        if not save_path:
            return

        try:
            wb = load_workbook(self.excel_path)
            ws = wb.active

            ws.cell(row=1, column=5, value="처리상태")
            ws.cell(row=1, column=6, value="상세결과")

            fill_map = {"yellow": YELLOW_FILL, "green": GREEN_FILL, "red": RED_FILL}
            for item in self.items:
                ws.cell(row=item.row_no, column=5, value=item.status)
                ws.cell(row=item.row_no, column=6, value=item.detail)
                if item.result_fill in fill_map:
                    ws[f"A{item.row_no}"].fill = fill_map[item.result_fill]

            wb.save(save_path)
            messagebox.showinfo("저장 완료", f"결과 파일 저장 완료\n{save_path}")
            self.log(f"[엑셀] 결과 저장 완료: {save_path}")
        except Exception as e:
            messagebox.showerror("저장 오류", f"결과 저장 실패\n{e}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
