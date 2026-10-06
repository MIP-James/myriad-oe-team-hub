# NaverAutomationLauncher.py

import os
import sys
import time
import threading
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext, simpledialog

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import SessionNotCreatedException
from selenium.webdriver.common.keys import Keys

# === Windows-safe path helpers ===
import re as _re_for_path
from pathlib import Path as _PathForPath

def _sanitize_filename(name: str, max_len: int = 120) -> str:
    # Remove forbidden chars on Windows and trim trailing dots/spaces
    name = name or ""
    name = _re_for_path.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', name)
    name = name.strip().strip('.')
    return name[:max_len] if max_len else name

def make_safe_save_path(base_dir: str, stem: str, ext: str = ".png") -> str:
    p = _PathForPath(base_dir).expanduser()
    try:
        p.mkdir(parents=True, exist_ok=True)
    except Exception:
        # If mkdir fails, fallback to user home
        p = _PathForPath.home() / "Downloads"
        try:
            p.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
    stem = _sanitize_filename(stem)
    out = p / f"{stem}{ext}"
    # Normalize to Windows style if on Windows
    try:
        return str(out.resolve())
    except Exception:
        return str(out)
# === end Windows-safe path helpers ===

from urllib.parse import urlparse

import random
import requests
import traceback
import tempfile
import datetime
from io import BytesIO
from PIL import Image as PilImage
from selenium.webdriver.chrome.options import Options
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage

# ================== AUTO.py 기능 ==================
def run_auto_py():
    def wait_for_user(msg):
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo("진행 안내", msg)
        root.destroy()

    folder = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(folder, "원본 이미지 URL.xlsx")
    pdf_path = os.path.join(folder, "사업자등록증미리어드아이피.pdf")
    poa_path = os.path.join(folder, "POA.pdf")
    idcard_path = os.path.join(folder, "신분증.pdf")

    def select_file(title):
        root = tk.Tk()
        root.withdraw()
        path = filedialog.askopenfilename(title=title, filetypes=[("모든 파일", "*.*")])
        root.destroy()
        return path

    if not os.path.exists(excel_path):
        excel_path = select_file("엑셀 파일을 선택하세요")
    if not os.path.exists(pdf_path):
        pdf_path = select_file("사업자등록증(증빙서류) 파일을 선택하세요")
    if not os.path.exists(poa_path):
        poa_path = select_file("POA 파일을 선택하세요")
    if not os.path.exists(idcard_path):
        idcard_path = select_file("신분증 파일을 선택하세요")

    # === 신고자 정보(이름/전화번호) 입력 팝업 ===
    def _prompt_name_phone():
        root = tk.Tk()
        root.withdraw()
        try:
            root.attributes('-topmost', True)
        except Exception:
            pass
        try:
            name = simpledialog.askstring("신고자 정보", "인증에 사용할 이름을 입력하세요 (빈칸 가능):", parent=root)
            phone = simpledialog.askstring("신고자 정보", "전화번호를 입력하세요 (하이픈 없이, 빈칸 가능):", parent=root)
        finally:
            try:
                root.destroy()
            except Exception:
                pass
        name = (name or "").strip()
        phone = (phone or "").strip()
        if phone and "-" in phone:
            messagebox.showinfo("전화번호 형식", "전화번호에 '-' 기호가 있어 제거했습니다.", parent=root)
            phone = phone.replace("-", "")
        return name, phone

    신고자_이름, 신고자_전화 = _prompt_name_phone()

    df = pd.read_excel(excel_path)
    상품_URL_list = df['상품 URL'].tolist()
    원본이미지_URL_list = df['원본이미지 (Link)'].tolist()
    증빙파일_경로 = pdf_path
    poa_file_경로 = poa_path
    idcard_file_경로 = idcard_path

    import webbrowser
    try:
        chrome_options = webdriver.ChromeOptions()
        chrome_options.add_experimental_option("detach", True)
        driver = webdriver.Chrome(options=chrome_options)
        driver.maximize_window()
    except SessionNotCreatedException:
        msg = (
            "Chrome 브라우저와 chromedriver 버전이 맞지 않아 자동화를 실행할 수 없습니다.\n"
            "크롬 브라우저 및 chromedriver를 같은 버전으로 맞춰주세요.\n"
            "다운로드 페이지를 열어드릴게요."
        )
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo("진행 안내", msg)
        root.destroy()
        webbrowser.open("https://chromedriver.chromium.org/downloads")
        sys.exit()
    except Exception as e:
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo("예상치 못한 오류", str(e))
        root.destroy()
        sys.exit()

    driver.get("https://right.naver.com")
    time.sleep(2)
    wait_for_user("네이버 로그인(수동) 후 '확인'을 눌러주세요.")
    wait = WebDriverWait(driver, 15)

    n = len(상품_URL_list)
    chunks = [ (상품_URL_list[i:i+10], 원본이미지_URL_list[i:i+10]) for i in range(0, n, 10)]

    for idx, (상품_URLs, 원본이미지_URLs) in enumerate(chunks):
        print(f"\n=== [{idx+1}/{len(chunks)}] {len(상품_URLs)}건 신고 시작 ===")
        wait.until(EC.element_to_be_clickable((By.XPATH, "//a[contains(., '신고하기')]"))).click()
        time.sleep(2)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//label[contains(., '네이버 쇼핑')]"))
        ).click()
        time.sleep(1)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//button[contains(., '신고하기') and not(@disabled)]"))
        ).click()
        time.sleep(2)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//label[contains(., '저작권')]"))
        ).click()
        time.sleep(2)
        if idx == 0:
            wait_for_user("개인정보 인증(본인인증 등)을 완료한 후, '확인'을 눌러주세요.")
        else:
            popup_report_btn_xpath = "//button[@type='button' and contains(@class,'Popup_btn_accept__ZiQkU') and text()='신고하기']"
            popup_report_btn = wait.until(EC.element_to_be_clickable((By.XPATH, popup_report_btn_xpath)))
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", popup_report_btn)
            time.sleep(0.5)
            try:
                popup_report_btn.click()
            except Exception as e:
                print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
                driver.execute_script("arguments[0].click();", popup_report_btn)
            time.sleep(1)
        wait = WebDriverWait(driver, 20)
        
        # === 신고자 정보: 이름/전화번호 자동 입력 (미입력 시 공란 유지) ===
        try:
            if 신고자_이름:
                try:
                    wait.until(EC.element_to_be_clickable((By.XPATH, "//input[@type='text' and contains(@placeholder, '이름')]"))).send_keys(신고자_이름)
                    time.sleep(0.3)
                except Exception:
                    pass
            if 신고자_전화:
                try:
                    wait.until(EC.element_to_be_clickable((By.XPATH, "//input[@type='text' and contains(@placeholder, '전화번호')]"))).send_keys(신고자_전화)
                    time.sleep(0.3)
                except Exception:
                    pass
        except Exception as _e:
            print("신고자(이름/전화) 자동 입력 실패:", _e)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//input[@placeholder='이메일을 입력해 주세요.']"))
        ).send_keys("onlinenforcement@myriadip.com")
        time.sleep(0.5)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//label[.//span[text()='안받음']]"))
        ).click()
        time.sleep(0.5)
        # --- '다음' 버튼 활성화 대기 후 클릭 (라벨 관련 로직 제거) ---
        CONTAINER_XPATH = "//*[@id='root']/div/div[1]/div/div/div/div/div/div[4]"
        NEXT_BTN_XPATH = f"{CONTAINER_XPATH}//button[contains(normalize-space(.), '다음')]"

        def _enabled_now(d):
            try:
                btn = d.find_element(By.XPATH, NEXT_BTN_XPATH)
                dis = btn.get_attribute("disabled")
                aria = (btn.get_attribute("aria-disabled") or "").lower() == "true"
                return btn.is_displayed() and (not dis) and (not aria)
            except Exception:
                return False

        WebDriverWait(driver, 10).until(lambda d: _enabled_now(d))
        time.sleep(0.3)

        try:
            btn = driver.find_element(By.XPATH, NEXT_BTN_XPATH)
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
            try:
                btn.click()
            except Exception:
                ActionChains(driver).move_to_element(btn).pause(0.05).click(btn).perform()
        except Exception:
            # 폴백: 강화된 클릭 로직
            click_next_strong(wait, driver, max_wait_seconds=8)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//button[contains(., '권리 추가 후 신고')] | //span[contains(., '권리 추가 후 신고')]/.."))
        ).click()
        time.sleep(2)
        file_input = wait.until(EC.presence_of_element_located((By.XPATH, "//input[@type='file']")))
        driver.execute_script("arguments[0].style.display = 'block';", file_input)
        file_input.send_keys(증빙파일_경로)
        time.sleep(2)
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//textarea[@placeholder='소유하고 계신 저작권에 대해 간단하게 설명해 주세요.']"))
        ).send_keys("원본 이미지 URL")
        time.sleep(0.5)
        url_add_btn_xpath = "//button[contains(., 'URL 입력란 추가')]"
        for _ in range(10):
            add_btn = WebDriverWait(driver, 10).until(
                EC.element_to_be_clickable((By.XPATH, url_add_btn_xpath))
            )
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", add_btn)
            time.sleep(0.2)
            try:
                add_btn.click()
            except Exception as e:
                print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
                driver.execute_script("arguments[0].click();", add_btn)
            time.sleep(0.2)
        URL_input_boxes = driver.find_elements(By.XPATH, "//input[@placeholder='URL 주소 입력']")
        for box, url in zip(URL_input_boxes, 원본이미지_URLs):
            box.send_keys(url)
            time.sleep(0.2)
        button = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//button[contains(., '다음')]")))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", button)
        time.sleep(1)
        try:
            button.click()
        except Exception as e:
            print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
            driver.execute_script("arguments[0].click();", button)
        time.sleep(2)
        wait.until(EC.element_to_be_clickable((By.XPATH, "//button[contains(., 'URL 신고대상 입력')]"))).click()
        time.sleep(1)
        for i, url in enumerate(상품_URLs):
            URL_box = wait.until(EC.element_to_be_clickable((By.XPATH, "//input[@placeholder='URL주소입력']")))
            URL_box.clear()
            URL_box.send_keys(url)
            time.sleep(1)
            driver.find_element(By.XPATH, "//button[span[text()='검색아이콘']]").click()
            time.sleep(2)
            desc_box = wait.until(EC.element_to_be_clickable(
                (By.XPATH, "//textarea[@placeholder=\"신고 대상을 확인할 수 있도록 ‘위치 또는 내용’ 을 기재해 주세요.\"]"))
            )
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", desc_box)
            time.sleep(0.5)
            try:
                desc_box.clear()
                desc_box.send_keys("이미지 도용")
            except Exception as e:
                print("기본 방식 실패, JS 클릭 후 send_keys 시도:", e)
                try:
                    driver.execute_script("arguments[0].click();", desc_box)
                    desc_box.clear()
                    desc_box.send_keys("이미지 도용")
                except Exception as e2:
                    print("JS 클릭도 실패:", e2)
            time.sleep(0.5)
            driver.find_element(By.XPATH, "//button[contains(., '신고대상 추가')]").click()
            time.sleep(1.2)
        wait.until(EC.element_to_be_clickable((By.XPATH, "//button[contains(., '확인')]"))).click()
        time.sleep(2)
        file_upload_input_xpath = "//input[@type='file' and @id='inpt_file']"
        file_input = wait.until(EC.presence_of_element_located((By.XPATH, file_upload_input_xpath)))
        driver.execute_script("arguments[0].style.display = 'block';", file_input)
        file_input.send_keys(poa_file_경로)
        time.sleep(1)
        file_input = driver.find_element(By.XPATH, file_upload_input_xpath)
        driver.execute_script("arguments[0].style.display = 'block';", file_input)
        file_input.send_keys(idcard_file_경로)
        time.sleep(1.5)
        agree_checkbox_xpath = "//label[span[contains(text(), '서약 및 동의')]]"
        agree_checkbox = wait.until(EC.element_to_be_clickable((By.XPATH, agree_checkbox_xpath)))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", agree_checkbox)
        try:
            agree_checkbox.click()
        except Exception as e:
            print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
            driver.execute_script("arguments[0].click();", agree_checkbox)
        time.sleep(0.5)
        report_btn_xpath = "//button[span[contains(text(), '신고하기')]] | //button[contains(text(), '신고하기')]"
        report_btn = wait.until(EC.element_to_be_clickable((By.XPATH, report_btn_xpath)))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", report_btn)
        time.sleep(0.5)
        try:
            report_btn.click()
        except Exception as e:
            print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
            driver.execute_script("arguments[0].click();", report_btn)
        time.sleep(2)
        # === '제출 전 확인' 페이지: '제출하기' 버튼 클릭 (2026-06-24 네이버 단계 추가) ===
        submit_confirm_xpath = (
            "//button[contains(., '제출하기')] | "
            "//*[@id='root']/div/div[1]/div/div/div/div/div/div[4]/ul/li[2]/button"
        )
        submit_confirm_btn = wait.until(EC.element_to_be_clickable((By.XPATH, submit_confirm_xpath)))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", submit_confirm_btn)
        time.sleep(0.5)
        try:
            submit_confirm_btn.click()
        except Exception as e:
            print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
            driver.execute_script("arguments[0].click();", submit_confirm_btn)
        time.sleep(2)
        home_btn_xpath = "//a[contains(@class, 'Header_lnk_home') and contains(text(), '권리보호센터')]"
        home_btn = wait.until(EC.element_to_be_clickable((By.XPATH, home_btn_xpath)))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", home_btn)
        time.sleep(0.5)
        try:
            home_btn.click()
        except Exception as e:
            print("기본 클릭 실패, JS로 강제 클릭 시도:", e)
            driver.execute_script("arguments[0].click();", home_btn)
        time.sleep(2)
        wait = WebDriverWait(driver, 15)

    print("\n=== 전체 자동화 완료! ===")
    wait_for_user("모든 자동 신고가 끝났습니다. 확인을 누르면 브라우저가 완전히 종료됩니다.")

    try:
        driver.quit()
    except Exception as e:
        print("드라이버 종료 중 에러:", e)

# ================== SmartStore_Crawler.py 기능 ==================
import os
import time
import random
import requests
import traceback
import tempfile
import datetime
from io import BytesIO
from PIL import Image as PilImage
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext

import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys


def _detect_chrome_major():
    """Registry 에서 설치된 Chrome major 버전 읽기. UC version_main=None 이 CfT 최신 드라이버를 받아오는 동작 회피용. 실패 시 None → UC 기본 동작."""
    try:
        import winreg
        for hive, path in (
            (winreg.HKEY_CURRENT_USER, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Google\Chrome\BLBeacon"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Wow6432Node\Google\Chrome\BLBeacon"),
        ):
            try:
                with winreg.OpenKey(hive, path) as key:
                    version, _ = winreg.QueryValueEx(key, "version")
                    return int(str(version).split(".")[0])
            except OSError:
                continue
    except Exception:
        pass
    return None


def start_hide_fixed_ui_elements_interval(driver, interval=200):
    js_code = f"""
    if (!window.__HIDE_UI_INTERVAL) {{
      window.__HIDE_UI_INTERVAL = setInterval(function() {{
        let selectors = [
          '#_productFloatingTab',
          '#_productFloatingTab > div'
        ];
        for(let s of selectors){{
          let elems = document.querySelectorAll(s);
          elems.forEach(e => {{
            e.style.display='none';
            e.style.visibility='hidden';
            e.style.opacity='0';
            e.style.height='0px';
            e.style.pointerEvents='none';
          }});    
        }}
      }}, {interval});
    }}
    """
    driver.execute_script(js_code)

def stop_hide_fixed_ui_elements_interval(driver):
    js_code = """
    if (window.__HIDE_UI_INTERVAL) {
        clearInterval(window.__HIDE_UI_INTERVAL);
        window.__HIDE_UI_INTERVAL = null;
    }
    """
    driver.execute_script(js_code)

def hide_scrollbar(driver):
    js_code = """
    document.body.style.overflow = 'hidden';
    document.documentElement.style.overflow = 'hidden';
    """
    driver.execute_script(js_code)

def show_scrollbar(driver):
    js_code = """
    document.body.style.overflow = '';
    document.documentElement.style.overflow = '';
    """
    driver.execute_script(js_code)

def capture_fullpage_by_scroll(driver, save_path, scroll_pause=0.5):
    viewport_height = driver.execute_script("return window.innerHeight")
    total_height = driver.execute_script("return document.body.scrollHeight")
    images = []
    driver.execute_script("window.scrollTo(0, 0)")
    time.sleep(scroll_pause)
    y = 0
    while y < total_height:
        driver.execute_script(f"window.scrollTo(0, {y});")
        time.sleep(scroll_pause)
        png = driver.get_screenshot_as_png()
        img = PilImage.open(BytesIO(png))
        images.append(img)
        y += viewport_height
    overlap = images[-1].height - (total_height % viewport_height or viewport_height)
    if overlap > 0 and len(images) > 1:
        images[-1] = images[-1].crop((0, overlap, images[-1].width, images[-1].height))
    stitched_height = sum(i.height for i in images)
    result_img = PilImage.new("RGB", (images[0].width, stitched_height))
    y_offset = 0
    for img in images:
        result_img.paste(img, (0, y_offset))
        y_offset += img.height
    result_img.save(save_path)

class SmartStoreCrawlerGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("네이버 스마트스토어 크롤러")
        self.root.geometry("670x520")
        self.create_widgets()
        self.excel_path = ""
        self.save_dir = ""
        self.driver = None
        self.wb = None
        self.ws = None
        self.image_refs = []
        self.should_stop = False

    def create_widgets(self):
        frm = tk.Frame(self.root)
        frm.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)
        row = 0
        tk.Label(frm, text="1. 엑셀 파일 선택:").grid(row=row, column=0, sticky='w')
        self.excel_entry = tk.Entry(frm, width=50)
        self.excel_entry.grid(row=row, column=1)
        tk.Button(frm, text="파일 찾기", command=self.select_excel).grid(row=row, column=2)
        row += 1
        tk.Label(frm, text="2. 이미지 저장 폴더:").grid(row=row, column=0, sticky='w')
        self.folder_entry = tk.Entry(frm, width=50)
        self.folder_entry.grid(row=row, column=1)
        tk.Button(frm, text="폴더 찾기", command=self.select_folder).grid(row=row, column=2)
        row += 1
        tk.Label(frm, text="3. User-Agent:").grid(row=row, column=0, sticky='w')
        self.ua_entry = tk.Entry(frm, width=50)
        self.ua_entry.insert(0, "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36")
        self.ua_entry.grid(row=row, column=1, columnspan=2, sticky='w')
        row += 1
        self.start_btn = tk.Button(frm, text="크롤링 시작", bg="#2DB400", fg="white", command=self.start_crawl)
        self.start_btn.grid(row=row, column=0, columnspan=3, sticky="ew", pady=7)
        row += 1
        self.progress = ttk.Progressbar(frm, orient="horizontal", length=500, mode="determinate")
        self.progress.grid(row=row, column=0, columnspan=3, pady=5)
        row += 1
        tk.Label(frm, text="로그:").grid(row=row, column=0, sticky='w')
        row += 1
        self.log = scrolledtext.ScrolledText(frm, width=90, height=15, state="normal")
        self.log.grid(row=row, column=0, columnspan=3)

    def select_excel(self):
        path = filedialog.askopenfilename(title="엑셀 파일 선택", filetypes=[("Excel files", "*.xlsx *.xlsm")])
        if path:
            self.excel_path = path
            self.excel_entry.delete(0, tk.END)
            self.excel_entry.insert(0, path)

    def select_folder(self):
        folder = filedialog.askdirectory(title="이미지 저장 폴더 선택")
        if folder:
            self.save_dir = folder
            self.folder_entry.delete(0, tk.END)
            self.folder_entry.insert(0, folder)

    def write_log(self, msg):
        self.log.insert(tk.END, msg)
        self.log.see(tk.END)
        self.root.update_idletasks()

    def start_crawl(self):
        self.start_btn["state"] = "disabled"
        self.log.delete(1.0, tk.END)
        threading.Thread(target=self.crawl_main, daemon=True).start()

    def is_service_block(self) -> bool:

        try:

            xpath_text = "//*[contains(normalize-space(.),'현재 서비스 접속이 불가합니다')]"

            if self.driver.find_elements(By.XPATH, xpath_text):

                return True

            frames = self.driver.find_elements(By.TAG_NAME, "iframe")[:3]

            for fr in frames:

                try:

                    self.driver.switch_to.frame(fr)

                    if self.driver.find_elements(By.XPATH, xpath_text):

                        return True

                finally:

                    self.driver.switch_to.default_content()

        except Exception:

            try:

                self.driver.switch_to.default_content()

            except Exception:

                pass

        return False
    def is_product_not_exist(self) -> bool:
        xp = "//*[contains(normalize-space(.),'상품이 존재하지 않습니다')]"
        try:
            # 현재 문서
            if self.driver.find_elements(By.XPATH, xp):
                return True
            # 상위 프레임 몇 개만 탐색
            frames = self.driver.find_elements(By.TAG_NAME, "iframe")[:5]
            for fr in frames:
                try:
                    self.driver.switch_to.frame(fr)
                    if self.driver.find_elements(By.XPATH, xp):
                        self.driver.switch_to.default_content()
                        return True
                finally:
                    try:
                        self.driver.switch_to.default_content()
                    except Exception:
                        pass
        except Exception:
            try:
                self.driver.switch_to.default_content()
            except Exception:
                pass
        return False





    def try_click_refresh_button(self, max_tries=3):
        for attempt in range(1, max_tries+1):
            try:
                xpath = "((//a|//button)[contains(normalize-space(.),'새로고침') or contains(@class,'refresh')])[1]"
                btns = self.driver.find_elements(By.XPATH, xpath)
                if btns:
                    self.write_log(f'새로고침 버튼 클릭 시도 ({attempt}/{max_tries})\n')
                    try:
                        btns[0].click()
                    except Exception:
                        self.driver.execute_script("arguments[0].click();", btns[0])
                    time.sleep(2.5)
                else:
                    self.write_log('새로고침 버튼 미탐지 → driver.refresh() 시도\n')
                    self.driver.refresh()
                    time.sleep(2.5)
                if not self.is_service_block():
                    self.write_log(f'새로고침({attempt}) 성공: 서비스불가 해제\n')
                    return True
                else:
                    self.write_log(f'새로고침({attempt}) 실패: 여전히 서비스불가\n')
            except Exception as e:
                self.write_log(f'새로고침 처리 예외({attempt}): {e}\n')
        return False
    def search_store_main_for_product(self, product_name):
        try:
            # 1) 검색 UI 열기(가능 시)
            try:
                search_btn = self.driver.find_element(
                    By.XPATH,
                    '//button[@data-shp-area="storh.search" or contains(@class,"storh")][contains(@class,"search") or contains(@aria-label,"검색")]'
                )
                try:
                    search_btn.click()
                except Exception:
                    self.driver.execute_script("arguments[0].click();", search_btn)
                time.sleep(0.5)
            except Exception:
                pass

            # 2) 인풋 찾기: placeholder 변형 허용 + 클래스 폴백
            input_xpath = (
                '//input['
                '@placeholder="검색어를 입력해주세요." or '
                '@placeholder="검색어를 입력해 주세요." or '
                '@placeholder="검색어를 입력해주세요" or '
                '@placeholder="검색어를 입력해 주세요"'
                ']'
            )
            inputs = self.driver.find_elements(By.XPATH, input_xpath)
            if not inputs:
                try:
                    inputs = [self.driver.find_element(By.XPATH, '//input[contains(@class,"_38DpYZHzTy")]')]
                except Exception:
                    pass
            if not inputs:
                self.write_log('스토어 메인 검색 인풋 탐색 실패\n')
                return

            box = inputs[0]
            try:
                box.clear()
            except Exception:
                pass
            box.click()
            box.send_keys(product_name)
            time.sleep(0.3)

            # 3) 돋보기 버튼 클릭 우선, 실패 시 Enter
            try:
                search_submit = self.driver.find_element(
                    By.XPATH, '//button[@data-shp-area="sch.search" or contains(@class,"search")]'
                )
                try:
                    search_submit.click()
                except Exception:
                    self.driver.execute_script("arguments[0].click();", search_submit)
            except Exception:
                box.send_keys(Keys.ENTER)

            time.sleep(1.5)

            # 4) 검색 결과 첫 상품 클릭 (가급적 /products/ 포함 링크)
            elems = self.driver.find_elements(By.XPATH, '(//a[contains(@href,"/products/")])[1]')
            if not elems:
                elems = self.driver.find_elements(By.CSS_SELECTOR, '#content a[href*="/products/"]')
            if elems:
                try:
                    self.driver.execute_script("arguments[0].click();", elems[0])
                except Exception:
                    elems[0].click()
                time.sleep(2.0)
                self.write_log('검색 결과 첫 상품 클릭 성공\n')
            else:
                self.write_log('검색 결과 게시물 클릭 실패: 요소 없음\n')
        except Exception as e:
            self.write_log(f'스토어 메인 검색/상품진입 실패: {e}\n')
            raise


    def is_security_challenge(self) -> bool:
        # 확장된 보안 배너/캡차 감지: message_wrap/message_text, <title>captcha</title>, ncpt captcha script, 일반 텍스트, iframe 스캔
        try:
            # 타이틀 기반 감지
            try:
                if (self.driver.title or '').strip().lower().find('captcha') != -1:
                    return True
            except Exception:
                pass

            # 스크립트 기반 (ncpt.naver.com ... captcha.js)
            script_hits = self.driver.find_elements(By.XPATH, "//script[contains(@src,'ncpt.naver.com') and contains(@src,'captcha')]")
            if script_hits:
                return True

            # DOM 기반 배너 감지
            xpath_candidates = [
                "//*[contains(@id,'message_text') and contains(@class,'message_text')]",
                "//*[contains(@id,'message_wrap') and contains(@class,'message_wrap')]",
                "//*[contains(normalize-space(.),'보안 확인을 완료해 주세요')]",
                "//*[contains(normalize-space(.),'보안 확인')]"
            ]
            for xp in xpath_candidates:
                if self.driver.find_elements(By.XPATH, xp):
                    return True

            # iframe 내부도 스캔 (최대 8개)
            frames = self.driver.find_elements(By.TAG_NAME, "iframe")[:8]
            for fr in frames:
                try:
                    self.driver.switch_to.frame(fr)
                    # 타이틀 기반
                    try:
                        if (self.driver.title or '').strip().lower().find('captcha') != -1:
                            return True
                    except Exception:
                        pass
                    # 스크립트/DOM 기반
                    if self.driver.find_elements(By.XPATH, "//script[contains(@src,'ncpt.naver.com') and contains(@src,'captcha')]"):
                        return True
                    for xp in xpath_candidates:
                        if self.driver.find_elements(By.XPATH, xp):
                            return True
                finally:
                    try:
                        self.driver.switch_to.default_content()
                    except Exception:
                        pass
        except Exception:
            try:
                self.driver.switch_to.default_content()
            except Exception:
                pass
        return False
    def check_security_challenge_pause(self):
        # 보안 확인 감지 시 팝업으로 작업자에게 안내하고, 해제될 때까지 대기
        try:
            while self.is_security_challenge():
                self.write_log("보안 확인 감지 → 작업자 인증 대기\n")
                # GUI가 멈추지 않도록 업데이트
                try:
                    self.root.update_idletasks()
                except Exception:
                    pass
                import tkinter.messagebox as _mb
                _mb.showinfo("보안 확인 필요", "브라우저에서 보안 확인(캡차/추가 인증)을 완료한 뒤\n'확인'을 눌러주세요.")
                # 잠시 대기 후 재검사
                import time as _t
                _t.sleep(0.8)
        except Exception as _e:
            try:
                self.write_log(f"보안 확인 감지 처리 예외: {_e}\n")
            except Exception:
                pass
    def wait_for_naver_login(self):
        self.write_log("네이버 로그인 페이지를 엽니다...\n")
        self.driver.get("https://nid.naver.com/nidlogin.login")
        messagebox.showinfo("네이버 로그인", "네이버에 로그인 한 뒤 '확인'을 눌러주세요.\n(로그인 후 이 창을 닫으세요)")
    
    def crawl_main(self):
        try:
            excel_path = self.excel_entry.get()
            save_dir = self.folder_entry.get()
            user_agent = self.ua_entry.get()

            if not os.path.exists(excel_path) or not os.path.exists(save_dir):
                self.write_log("엑셀 파일/폴더를 모두 선택하세요.\n")
                self.start_btn["state"] = "normal"
                return

            self.wb = load_workbook(excel_path)
            self.ws = self.wb.active

            url_list = []
            for idx, row in enumerate(self.ws.iter_rows(min_row=2, min_col=7, max_col=7), start=2):
                cell = row[0]
                url = None
                if cell.hyperlink and cell.hyperlink.target and cell.hyperlink.target.strip().startswith("http"):
                    url = cell.hyperlink.target.strip()
                elif cell.value and isinstance(cell.value, str) and cell.value.strip().startswith("http"):
                    url = cell.value.strip()
                if url:
                    url_list.append((idx, url))
            self.write_log(f"총 {len(url_list)}개의 유효한 URL을 읽었습니다.\n")
            self.progress["maximum"] = len(url_list)

            chrome_options = uc.ChromeOptions()
            chrome_options.add_argument(f'--user-agent={user_agent}')
            chrome_options.add_argument("--disable-logging")
            # chrome_options.add_argument("--headless")  # 필요시 사용
            self.driver = uc.Chrome(options=chrome_options, use_subprocess=True, version_main=_detect_chrome_major())  # frozen launcher 환경에서 통합 런처 재실행 방지 + CfT 최신 드라이버 자동 다운로드 회피

            # 로그인 먼저 진행!
            self.wait_for_naver_login()

            self.driver.minimize_window()
            batch_size = 10

            for i in range(0, len(url_list), batch_size):
                batch = url_list[i:i + batch_size]
                tab_mapping = []

                # 첫 번째 URL (기존 탭)
                row_num, url = batch[0]
                self.driver.get(url)
                # 보안 확인 감지 시 작업자 인증 대기
                self.check_security_challenge_pause()
                tab_mapping.append((self.driver.current_window_handle, row_num, url))
                time.sleep(random.uniform(1.0, 1.5))

                # 나머지 URL (새 탭)
                for row_num, url in batch[1:]:
                    self.driver.switch_to.new_window('tab')
                    self.driver.get(url)
                    # 보안 확인 감지 시 작업자 인증 대기
                    self.check_security_challenge_pause()
                    tab_mapping.append((self.driver.current_window_handle, row_num, url))
                    time.sleep(random.uniform(1.0, 1.5))

                first_cycle = True
                last_tab_had_security = True
                prev_url_for_domain = None
                for handle, row_num, url in tab_mapping:
                    try:
                        # ✅ 1) 반드시 해당 핸들로 스위치
                        try:
                            self.driver.switch_to.window(handle)
                        except Exception:
                            self.driver.switch_to.new_window('tab')
                            self.driver.get(url)
                            handle = self.driver.current_window_handle
                        time.sleep(0.15 if first_cycle else 0.07)

                        # ✅ 2) 보안 확인은 스위치 이후에 체크
                        same_domain = False
                        try:
                            curu = (self.driver.current_url or "")
                            if prev_url_for_domain and curu:
                                same_domain = urlparse(prev_url_for_domain).netloc == urlparse(curu).netloc
                        except Exception:
                            pass
                        if (not last_tab_had_security) and same_domain:
                            pass
                        else:
                            self.check_security_challenge_pause()
                        product_name = self.ws[f'J{row_num}'].value or ""
                        first_cycle = False

                        service_block = ([1] if self.is_service_block() else [])
                        last_tab_had_security = bool(service_block)
                        try:
                            prev_url_for_domain = (self.driver.current_url or prev_url_for_domain)
                        except Exception:
                            pass
                        if service_block:
                            self.write_log(f'[{row_num}] 서비스불가 감지 - 새로고침 최대 3회 시도\n')
                            refresh_success = self.try_click_refresh_button(max_tries=3)
                            if refresh_success:
                                try:
                                    self.write_log(f'[{row_num}] 새로고침 성공, 스토어 메인 검색 로직 진입\n')
                                    self.search_store_main_for_product(product_name)
                                except Exception as e:
                                    self.write_log(f'[{row_num}] 스토어 메인 검색 우회 실패 → N열 Store URL 우회로 진입\n')
                                    service_block = [1]
                            else:
                                self.write_log(f'[{row_num}] 새로고침 3회 실패, N열 Store URL 우회로 진입\n')

                        service_block = ([1] if self.is_service_block() else [])
                        if service_block:
                            n_cell = self.ws[f'N{row_num}']
                            n_url = None
                            if n_cell.hyperlink and n_cell.hyperlink.target and n_cell.hyperlink.target.strip().startswith("http"):
                                n_url = n_cell.hyperlink.target.strip()
                            elif n_cell.value and isinstance(n_cell.value, str) and n_cell.value.strip().startswith("http"):
                                n_url = n_cell.value.strip()
                            if n_url:
                                self.driver.get(n_url)
                                time.sleep(2)
                                try:
                                    self.search_store_main_for_product(product_name)
                                except Exception as e:
                                    self.write_log(f'[{row_num}] N열 Store URL 검색 우회 실패: {e}\n')
                                    continue

                        
                        if self.is_product_not_exist():
                            self.write_log(f'[{row_num}] 상품이 존재하지 않습니다. → 새로고침 최대 3회 시도\n')
                            refreshed = False
                            for attempt in range(1, 4):
                                try:
                                    self.driver.refresh()
                                    time.sleep(2.5)
                                    if not self.is_product_not_exist():
                                        self.write_log(f'새로고침({attempt}) 성공: 상품 페이지 복구\n')
                                        refreshed = True
                                        break
                                    else:
                                        self.write_log(f'새로고침({attempt}) 실패: 여전히 미노출\n')
                                except Exception as e:
                                    self.write_log(f'새로고침 예외({attempt}): {e}\n')
                            if not refreshed:
                                self.write_log(f'[{row_num}] 새로고침 3회 실패 → 스토어 검색 우회 진입\n')
                                try:
                                    self.search_store_main_for_product(product_name)
                                except Exception as e:
                                    self.write_log(f'[{row_num}] 검색 결과 게시물 클릭 실패: {e}\n')

                        try:
                            img_url = ""

                            try:
                                # Updated: robust representative image selection with multiple fallbacks
                                img_url = ""
                                img_elem = None
                                # 1) Current CSS path observed (container > ... > img)
                                try:
                                    img_elem = self.driver.find_element(By.CSS_SELECTOR, '#content > div > div.Cpf2P_YsRS > div.OaKLUocIcJ.Ux1iE3sNAr > div.PYE1T66W79 > div.mdFeBiFowv > img')
                                except Exception:
                                    pass
                                # 2) Class + alt attribute
                                if img_elem is None:
                                    try:
                                        img_elem = self.driver.find_element(By.CSS_SELECTOR, 'img.TgO1N1wWTm[alt="대표이미지"]')
                                    except Exception:
                                        pass
                                # 3) Provided absolute XPath
                                if img_elem is None:
                                    try:
                                        img_elem = self.driver.find_element(By.XPATH, '//*[@id="content"]/div/div[2]/div[1]/div[1]/div[2]/img')
                                    except Exception:
                                        pass
                                # 4) Generic alt-based XPath
                                if img_elem is None:
                                    try:
                                        img_elem = self.driver.find_element(By.XPATH, '//img[@alt="대표이미지"]')
                                    except Exception:
                                        pass
                                # 5) Legacy selector (backward compatibility)
                                if img_elem is None:
                                    try:
                                        img_elem = self.driver.find_element(By.CSS_SELECTOR, 'div.bd_1uFKu img.bd_2DO68')
                                    except Exception:
                                        pass

                                if img_elem is not None:
                                    # Prefer src, then data-src, then srcset (first URL)
                                    img_url = (img_elem.get_attribute('src') or
                                               img_elem.get_attribute('data-src') or
                                               '')
                                    if not img_url:
                                        srcset = img_elem.get_attribute('srcset') or ''
                                        if srcset:
                                            # Take the first URL before whitespace
                                            img_url = srcset.split(',')[0].strip().split(' ')[0]
                                else:
                                    raise Exception('대표 이미지 요소를 찾을 수 없음')
                            except Exception:
                                self.write_log(f"[{row_num}] 대표 이미지 로드 실패\n")

                            if img_url:
                                try:
                                    img_data = requests.get(img_url, timeout=10).content
                                    pil_img = PilImage.open(BytesIO(img_data)).convert("RGBA").resize((115, 137))
                                    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
                                    pil_img.save(temp_file.name, format="PNG")
                                    img = OpenpyxlImage(temp_file.name)
                                    img.anchor = f'F{row_num}'
                                    self.ws.add_image(img)
                                    self.image_refs.append((temp_file.name, img))
                                    self.write_log(f"[{row_num}] 대표 이미지 엑셀 업로드 완료\n")
                                except Exception as e:
                                    self.write_log(f"[{row_num}] 이미지 처리 오류: {e}\n")

                            price = ""
                            try:
                                price_elem = self.driver.find_element(
                                    By.XPATH,
                                    '//strong[contains(@class,"Xu9MEKUuIo") and contains(@class,"s6EKUu28OE")]/span[contains(@class,"e1DMQNBPJ_")]'
                                )
                                price = price_elem.text
                            except Exception:
                                self.write_log(f"[{row_num}] 가격 정보 로드 실패\n")
                            self.ws[f'L{row_num}'] = price

                            # === Store Name (M column) extraction & write ===
                            try:
                                wait_tmp = WebDriverWait(self.driver, 2)
                                store_name = ""
                                xpaths = [
                                    "//h1[contains(@class,'ic78vHV230')]//div[@id='pc-gnbWidget']//a[contains(@class,'_nlog_click')]//span[contains(@class,'rj3UIhnh81')]",
                                    "//*[@id='pc-gnbWidget']//a[contains(@class,'_nlog_click')]//span",
                                    "//*[@id='pc-gnbWidget']//*[contains(@data-shp-area,'storh.home')]//span[1]",
                                    "//h1//span[1]",
                                ]
                                for xp in xpaths:
                                    try:
                                        el = wait_tmp.until(EC.presence_of_element_located((By.XPATH, xp)))
                                        try:
                                            wait_tmp.until(EC.visibility_of(el))
                                        except Exception:
                                            pass
                                        txt = (el.text or "").strip()
                                        if txt:
                                            store_name = txt
                                            break
                                    except Exception:
                                        continue
                                try:
                                    try:
                                        self._store_name_cache
                                    except Exception:
                                        self._store_name_cache = {}
                                    try:
                                        cur_domain = urlparse((self.driver.current_url or '')).netloc
                                    except Exception:
                                        cur_domain = ''
                                    if not store_name and cur_domain in self._store_name_cache:
                                        store_name = self._store_name_cache.get(cur_domain, '')
                                    elif store_name and cur_domain:
                                        self._store_name_cache[cur_domain] = store_name
                                    self.ws[f'M{row_num}'] = store_name
                                except Exception:
                                    pass
                            except Exception:
                                pass
                            # === end Store Name extraction ===


                            self.write_log(f"[{row_num}] 가격 정보 엑셀 업로드 완료\n")

                            detail_btn_clicked = False
                            for _ in range(15):
                                try:
                                    btn = self.driver.find_element(
                                        By.XPATH,
                                        '//button[contains(@class,"dnjZ5j9eOR") and contains(@class,"_nlog_click") and text()="상세정보 펼쳐보기"]'
                                    )
                                    if btn.is_displayed():
                                        self.driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
                                        time.sleep(0.5)
                                        btn.click()
                                        detail_btn_clicked = True
                                        self.write_log(f'[{row_num}] 상세정보 펼쳐보기 클릭 성공!\n')
                                        break
                                    else:
                                        self.driver.find_element(By.TAG_NAME, 'body').send_keys(Keys.END)
                                        time.sleep(0.2)
                                except Exception:
                                    self.driver.find_element(By.TAG_NAME, 'body').send_keys(Keys.END)
                                    time.sleep(0.2)
                            if not detail_btn_clicked:
                                self.write_log(f'[{row_num}] 상세정보 펼쳐보기 버튼 없음!\n')

                            time.sleep(2.5)

                            try:
                                end_time = time.time() + 3
                                while time.time() < end_time:
                                    self.driver.find_element(By.TAG_NAME, 'body').send_keys(Keys.END)
                                    time.sleep(0.2)
                                self.write_log(f'[{row_num}] 3초간 최하단 스크롤 완료\n')
                            except Exception as e:
                                self.write_log(f'[{row_num}] 스크롤 다운 오류: {e}\n')

                            try:
                                start_hide_fixed_ui_elements_interval(self.driver, interval=200)
                                hide_scrollbar(self.driver)
                                time.sleep(0.6)
                                today = datetime.datetime.now().strftime("%Y%m%d")
                                product_code = self.ws[f'H{row_num}'].value or "NO_CODE"
                                file_stem = f"{product_code}_{today}"
                                save_path = make_safe_save_path(self.save_dir, file_stem, ext=".png")
                                capture_fullpage_by_scroll(self.driver, save_path, scroll_pause=0.7)
                                stop_hide_fixed_ui_elements_interval(self.driver)
                                show_scrollbar(self.driver)
                                self.write_log(f'[{row_num}] GoFullPage 스타일(setInterval+스크롤바) 캡처 저장: {save_path}\n')
                            except Exception as e:
                                self.write_log(f'[{row_num}] 뷰포트 캡처 오류: {e}\n')

                            self.write_log(f"[{row_num}] 크롤링 완료 | 가격: {price}\n")
                        except Exception as crawl_error:
                            self.write_log(f"[{row_num}] 크롤링 중 오류: {crawl_error}\n")
                            traceback.print_exc()
                            continue

                    except Exception as e:
                        self.write_log(f"[{row_num}] 탭 처리 중 예외: {e}\n")
                        traceback.print_exc()

                    self.progress["value"] += 1
                    self.root.update_idletasks()

                try:
                    self.wb.save(self.excel_path)
                    self.write_log("엑셀 저장 완료!\n")
                except Exception as e:
                    self.write_log(f"엑셀 저장 실패: {e}\n")
                    traceback.print_exc()

                for handle, _, _ in tab_mapping[1:]:
                    try:
                        # ✅ 닫기 전에 해당 탭으로 스위치
                        self.driver.switch_to.window(handle)
                        self.check_security_challenge_pause()
                        self.driver.close()
                        time.sleep(0.1)
                    except:
                        pass
                # ✅ 마지막에 첫 탭으로 복귀
                try:
                    self.driver.switch_to.window(tab_mapping[0][0])
                except Exception:
                    try:
                        remaining = self.driver.window_handles
                        if remaining:
                            self.driver.switch_to.window(remaining[0])
                    except Exception:
                        pass

            messagebox.showinfo("작업 완료", "모든 URL 크롤링 및 엑셀 저장이 완료되었습니다!")

        except Exception as e:
            err_msg = str(e)
            if "only supports Chrome version" in err_msg or "SessionNotCreatedException" in err_msg or "DevToolsActivePort file doesn't exist" in err_msg:
                version_msg = (
                    "Chrome과 chromedriver 버전이 맞지 않습니다!\n"
                    "최신 크롬 또는 chromedriver를 설치하세요.\n"
                    "https://chromedriver.chromium.org/downloads 에서 받으실 수 있습니다."
                )
                self.write_log(version_msg + "\n")
                messagebox.showerror("ChromeDriver 버전 불일치", version_msg)
            else:
                self.write_log(f"전체 처리 중 예외: {e}\n")
            traceback.print_exc()
        finally:
            try:
                if self.wb:
                    self.wb.save(self.excel_path)
            except Exception as e:
                self.write_log(f"마지막 save 오류: {e}\n")
            try:
                if self.wb:
                    self.wb.close()
            except Exception as e:
                self.write_log(f"마지막 close 오류: {e}\n")
            try:
                if self.driver:
                    self.driver.quit()
            except Exception as e:
                self.write_log(f"브라우저 종료 오류: {e}\n")
            for fname, _ in getattr(self, "image_refs", []):
                try:
                    os.remove(fname)
                except Exception:
                    pass
            self.write_log("브라우저 종료 및 정리 완료!\n")
            self.start_btn["state"] = "normal"
            self.root.quit()

def run_smartstore_py():
    root = tk.Tk()
    app = SmartStoreCrawlerGUI(root)
    root.mainloop()

# ================== 런처 메인 ==================
def start_auto():
    threading.Thread(target=run_auto_py, daemon=True).start()

def start_smartstore():
    threading.Thread(target=run_smartstore_py, daemon=True).start()

if __name__ == "__main__":
    root = tk.Tk()
    root.title("네이버 자동화 통합 런처")
    root.geometry("400x200")
    root.resizable(False, False)
    tk.Label(root, text="실행할 자동화 프로그램을 선택하세요.", font=("맑은 고딕", 13)).pack(pady=20)
    tk.Button(root, text="스마트스토어 상품 크롤러 실행", height=2, width=25, command=start_smartstore).pack(pady=10)
    tk.Button(root, text="저작권 신고 자동화 실행", height=2, width=25, command=start_auto).pack(pady=5)
    root.mainloop()
from selenium.webdriver.common.action_chains import ActionChains

from selenium.webdriver.common.action_chains import ActionChains

def _find_clickable_next_buttons(driver):
    """여러 셀렉터로 '다음' 버튼 후보를 수집하고, 화면에 보이고 enable 된 것만 반환"""
    XPATHS = [
        "//button[normalize-space()='다음' and not(@disabled) and not(@aria-disabled='true')]",
        "//button[.//span[normalize-space()='다음'] and not(@disabled) and not(@aria-disabled='true')]",
        "//button[contains(normalize-space(.), '다음') and not(@disabled) and not(@aria-disabled='true')]",
        "//*[@id='root']/div/div[1]/div/div/div/div/div/div[4]/ul/li/button",
    ]
    CSSS = [
        "#root > div > div.Layout_container__3lZO0 > div > div > div > div > div > div.Reportinner_form_btn__CxRob > ul > li > button",
        r"button.Reportinner_btn__HNWT\+",
        "button.Reportinner_btn_lg__M-fNV",
        "button.Reportinner_btn_color02__u2Qcx",
    ]

    candidates = []
    for xp in XPATHS:
        try:
            candidates.extend(driver.find_elements(By.XPATH, xp))
        except Exception:
            pass
    for cs in CSSS:
        try:
            candidates.extend(driver.find_elements(By.CSS_SELECTOR, cs))
        except Exception:
            pass

    filtered = []
    for el in candidates:
        try:
            if el.is_displayed():
                aria_dis = (el.get_attribute("aria-disabled") or "").lower() == "true"
                dis = el.get_attribute("disabled")
                if not dis and not aria_dis:
                    filtered.append(el)
        except Exception:
            continue

    seen = set()
    uniq = []
    for el in filtered:
        try:
            key = el.id
        except Exception:
            key = id(el)
        if key not in seen:
            seen.add(key)
            uniq.append(el)
    return uniq

def click_next_strong(wait, driver, max_wait_seconds=12):
    """'다음' 버튼 클릭 강화판: enable 대기 → 다중 셀렉터 → JS(React 친화) → 좌표 클릭 → iframe(재귀) → 최종 폴백"""
    import time
    end_time = time.time() + max_wait_seconds

    def _visible(el):
        try:
            if not el.is_displayed():
                return False
            box = el.rect if hasattr(el, "rect") else None
            if box and (box.get("width", 1) == 0 or box.get("height", 1) == 0):
                return False
            return True
        except Exception:
            return False

    def _js_click_by_text():
        # React에서도 잘 먹히는 디스패치 방식 + Shadow DOM 탐색 포함
        script = r"""
        (function() {
          function normalize(t){return (t||'').replace(/\s+/g,' ').trim();}
          function isEnabled(el){
            if(!el) return false;
            if(el.disabled) return false;
            const a = el.getAttribute('aria-disabled');
            if(a && a.toLowerCase()==='true') return false;
            return true;
          }
          function isShown(el){
            if(!el) return false;
            const cs = getComputedStyle(el);
            if(cs.display==='none' || cs.visibility==='hidden' || cs.opacity==='0') return false;
            const r = el.getBoundingClientRect();
            if(r.width===0 || r.height===0) return false;
            return true;
          }
          function collect(root){
            const out = [];
            const walker = (node)=>{
              try{
                out.push(...node.querySelectorAll('button,[role="button"]'));
                const kids = node.querySelectorAll('*');
                kids.forEach(k => { if(k.shadowRoot){ walker(k.shadowRoot); } });
              }catch(e){}
            };
            walker(root);
            return out;
          }
          const all = collect(document);
          const cand = all.filter(b => normalize(b.textContent).includes('다음'));
          let btn = null;
          for(const b of cand){
            if(isEnabled(b) && isShown(b)){ btn = b; break; }
          }
          if(!btn && cand.length>0){ btn = cand[0]; }
          if(!btn) return false;

          try { btn.scrollIntoView({block:'center'}); } catch(e){}

          const types = ['mouseover','pointerover','mousedown','pointerdown','mouseup','pointerup','click'];
          try{
            for(const t of types){
              const ev = new MouseEvent(t, {bubbles:true, cancelable:true, view:window});
              btn.dispatchEvent(ev);
            }
          }catch(e){
            try{ btn.click(); }catch(e2){}
          }
          return true;
        })();
        """
        try:
            return bool(driver.execute_script(script))
        except Exception:
            return False

    def _actions_click(el):
        try:
            from selenium.webdriver.common.action_chains import ActionChains
            r = el.rect
            x_off = int((r.get("width", 0))/2) - 1
            y_off = int((r.get("height", 0))/2) - 1
            ActionChains(driver).move_to_element_with_offset(el, x_off, y_off).click().perform()
            return True
        except Exception:
            try:
                el.click()
                return True
            except Exception:
                try:
                    driver.execute_script("arguments[0].click();", el)
                    return True
                except Exception:
                    return False

    def _collect_candidates():
        XPATHS = [
            "//button[normalize-space()='다음' and not(@disabled) and not(@aria-disabled='true')]",
            "//button[.//span[normalize-space()='다음'] and not(@disabled) and not(@aria-disabled='true')]",
            "//button[contains(normalize-space(.), '다음') and not(@disabled) and not(@aria-disabled='true')]",
            "//*[@id='root']/div/div[1]/div/div/div/div/div/div[4]/ul/li/button",
        ]
        CSSS = [
            "#root > div > div.Layout_container__3lZO0 > div > div > div > div > div > div.Reportinner_form_btn__CxRob > ul > li > button",
            r"button.Reportinner_btn__HNWT\+",
            "button.Reportinner_btn_lg__M-fNV",
            "button.Reportinner_btn_color02__u2Qcx",
        ]
        cand = []
        for xp in XPATHS:
            try: cand.extend(driver.find_elements(By.XPATH, xp))
            except Exception: pass
        for cs in CSSS:
            try: cand.extend(driver.find_elements(By.CSS_SELECTOR, cs))
            except Exception: pass
        out = []
        seen = set()
        for el in cand:
            try:
                key = el.id
            except Exception:
                key = id(el)
            if key in seen: continue
            seen.add(key)
            if _visible(el): out.append(el)
        return out

    def _poll_enable_and_click(timeout_sec=5):
        t_end = time.time() + timeout_sec
        while time.time() < t_end:
            try:
                btn = wait.until(EC.element_to_be_clickable((
                    By.XPATH, "//button[(normalize-space()='다음' or .//span[normalize-space()='다음'] or contains(normalize-space(.),'다음')) and not(@disabled) and not(@aria-disabled='true')]"
                )))
                if _actions_click(btn):
                    return True
            except Exception:
                pass
            time.sleep(0.25)
        return False

    if _poll_enable_and_click(5):
        return True

    for el in _collect_candidates():
        if _actions_click(el):
            return True

    if _js_click_by_text():
        return True

    def _try_iframes(depth=0, max_depth=2):
        if depth > max_depth: return False
        frames = driver.find_elements(By.TAG_NAME, "iframe")[:10]
        for fr in frames:
            try:
                driver.switch_to.frame(fr)
                if _poll_enable_and_click(2):
                    driver.switch_to.default_content()
                    return True
                for el in _collect_candidates():
                    if _actions_click(el):
                        driver.switch_to.default_content()
                        return True
                if _js_click_by_text():
                    driver.switch_to.default_content()
                    return True
                if _try_iframes(depth+1, max_depth):
                    driver.switch_to.default_content()
                    return True
            except Exception:
                pass
            finally:
                try: driver.switch_to.default_content()
                except Exception: pass
        return False

    if _try_iframes():
        return True

    try:
        from selenium.webdriver.common.keys import Keys
        driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ENTER)
        time.sleep(0.5)
    except Exception:
        pass

    return False

def click_next_via_elementfrompoint(driver):
    script = r"""
    (function(){
      function q(sel, root=document){ try { return root.querySelector(sel); } catch(e){ return null; } }
      function all(sel, root=document){ try { return Array.from(root.querySelectorAll(sel)); } catch(e){ return []; } }
      function norm(t){ return (t||'').replace(/\s+/g,' ').trim(); }
      var container = q("#root > div > div.Layout_container__3lZO0 > div > div > div > div > div > div.Reportinner_form_btn__CxRob")
                   || q("div.Reportinner_form_btn__CxRob");
      if(!container) return false;
      var btn = q("button", container);
      if(!btn || norm(btn.textContent).indexOf("다음") === -1) {
        btn = all("button", container).find(b=> norm(b.textContent).indexOf("다음")>-1);
        if(!btn) return false;
      }
      try { btn.scrollIntoView({block:'center'}); } catch(e){}
      var r = btn.getBoundingClientRect();
      var cx = Math.floor(r.left + r.width/2);
      var cy = Math.floor(r.top + r.height/2);
      var cleaned = [];
      for (var tries=0; tries<6; tries++){
        var topEl = document.elementFromPoint(cx, cy);
        if(!topEl) break;
        if (topEl === btn || btn.contains(topEl)) break;
        var cs = getComputedStyle(topEl);
        var isCover = (cs.position === 'fixed' || cs.position === 'sticky' || cs.position === 'absolute');
        if(isCover){
          topEl.setAttribute('data-pe', topEl.style.pointerEvents || '');
          topEl.style.pointerEvents = 'none';
          cleaned.push(topEl);
        } else {
          break;
        }
      }
      var types = ['mouseover','pointerover','mousedown','pointerdown','mouseup','pointerup','click'];
      try { for (var i=0;i<types.length;i++){ var ev = new MouseEvent(types[i], {bubbles:true, cancelable:true, view:window}); btn.dispatchEvent(ev);} }
      catch(e){ try { btn.click(); } catch(e2){} }
      cleaned.forEach(el=>{ var prev = el.getAttribute('data-pe'); el.style.pointerEvents = prev; el.removeAttribute('data-pe'); });
      return true;
    })();
    """
    try:
        return bool(driver.execute_script(script))
    except Exception:
        return False


def submit_next_form(driver):
    script = r"""
    (function(){
      function norm(t){ return (t||'').replace(/\s+/g,' ').trim(); }
      var container = document.querySelector("#root > div > div.Layout_container__3lZO0 > div > div > div > div > div > div.Reportinner_form_btn__CxRob")
                   || document.querySelector("div.Reportinner_form_btn__CxRob");
      if(!container) return false;
      var btn = container.querySelector("button");
      if(!btn || norm(btn.textContent).indexOf("다음")===-1){
        var all = container.querySelectorAll("button");
        for (var i=0;i<all.length;i++){
          if(norm(all[i].textContent).indexOf("다음")>-1){ btn = all[i]; break; }
        }
      }
      if(!btn) return false;
      var form = btn.closest('form');
      if(form){
        if (typeof form.requestSubmit === 'function') form.requestSubmit();
        else form.submit();
        return true;
      }
      try { btn.click(); return true; } catch(e){}
      return false;
    })();
    """
    try:
        return bool(driver.execute_script(script))
    except Exception:
        return False


def focus_and_press_enter(driver):
    script = r"""
    (function(){
      function norm(t){ return (t||'').replace(/\s+/g,' ').trim(); }
      var container = document.querySelector("#root > div > div.Layout_container__3lZO0 > div > div > div > div > div > div.Reportinner_form_btn__CxRob")
                   || document.querySelector("div.Reportinner_form_btn__CxRob");
      if(!container) return false;
      var btn = Array.from(container.querySelectorAll("button")).find(b=>norm(b.textContent).includes("다음"));
      if(!btn) return false;
      try{ btn.scrollIntoView({block:'center'}); }catch(e){}
      try{ btn.focus(); }catch(e){}
      var ev = new KeyboardEvent('keydown', {key:'Enter', code:'Enter', keyCode:13, which:13, bubbles:true});
      btn.dispatchEvent(ev);
      var ev2 = new KeyboardEvent('keyup', {key:'Enter', code:'Enter', keyCode:13, which:13, bubbles:true});
      btn.dispatchEvent(ev2);
      try{ btn.click(); }catch(e){}
      return true;
    })();
    """
    try:
        return bool(driver.execute_script(script))
    except Exception:
        return False


def ensure_top_context(driver):
    try:
        driver.switch_to.default_content()
    except Exception:
        pass
    try:
        handles = driver.window_handles
        if handles and driver.current_window_handle != handles[-1]:
            driver.switch_to.window(handles[-1])
    except Exception:
        pass